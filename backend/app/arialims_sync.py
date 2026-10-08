"""AriaLIMS → analysis_samples sync.

Each configured sampling point (``AriaLimsSamplingPoint``, keyed by SCID) is requested from
``/api/AriaLIMS/results``. The API returns one row per measured parameter; rows of one point
that share a ``samplingTime`` become one ``AnalysisSample`` whose ``parameters`` hold
``{analysis name (or mapped key): value}``.
"""
from __future__ import annotations

import logging
import re
from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable

from sqlalchemy.orm import Session

from . import models
from .arialims_client import (
    ANALYSIS_TYPES,
    AriaLimsConfig,
    AriaLimsNotConfigured,
    endpoint_catalog,
    extract_results,
    fetch_results,
    normalize_base_url,
    test_connection,
)

logger = logging.getLogger("pvc_arvand.arialims")

ProgressCb = Callable[[int, int], None] | None

# Access Analyse scopes a sampling point can feed.
VALID_SCOPES = {"element", "group", "sub_plant", "electrolyzer", "total_plant"}

_NUMBER = re.compile(r"^[+-]?\d+(?:[.,]\d+)?(?:[eE][+-]?\d+)?$")

# AriaLIMS analysis name (lower case) → parameter key of the Analysis forms (per-electrolyzer
# / per-cell / per-train / per-group forms; see routers/analyses.py). A point's own
# ``parameter_map`` overrides these. Names not listed are stored under the AriaLIMS name.
DEFAULT_PARAMETER_MAPS: dict[str, dict[str, str]] = {
    "anolyte": {
        "ph": "pH An",
        "nacl": "NaCl An",
        "naclo3": "NaClO3 An",
        "na2so4": "Na2SO4 An",
        "naocl": "NaOCl An",
        "hcl": "HCl An",
        "density at 20'c": "density at 20° An",
        "temperature": "temperature An",
    },
    "catholyte": {
        "naoh": "NaOH",
        "nacl": "NaCl",
        "naclo3": "NaClO3",
        "na2so4": "Na2SO4",
        "fe": "Fe",
        "temperature": "temperature",
    },
    "pure_brine": {
        "ph": "pH",
        "density at 20'c": "density at 20 °C",
        "nacl": "NaCl Pb",
        "naclo3": "NaClO3 Pb",
        "na2co3": "Na2CO3 Pb",
        "naoh": "NaOH Pb",
        "na2so4": "Na2SO4 Pb",
        "naocl": "NaOCl Pb",
        "hcl": "HCl Pb",
        "ca+mg": "Ca +Mg",
        "sio2": "SiO2",
        "toc": "Organics",
    },
    "chlorine_gas": {
        "cl2 (dry bas.)": "Cl2",
        "co2": "CO2",
        "rest gas": "restgas Cl",
        "o2 (dry bas.)": "O2 Cl",
        "h2": "H2 Cl",
        "n2": "N2 Cl",
        "br": "Br Cl",
    },
    "hcl": {"hcl": "HCl", "density at 20'c": "Dichte"},
}

CL2_CO2_KEY = "Cl2 + CO2 Cl"
BATCH_SIZE = 30  # SCIDs per request

# "UnitTag" column of the AriaLIMS sample-point list → analysis type.
UNIT_TAG_TYPES = (
    ("pure brine", "pure_brine"),
    ("anolyte", "anolyte"),
    ("catholyte", "catholyte"),
    ("cl2", "chlorine_gas"),
    ("chlorine", "chlorine_gas"),
    ("hcl", "hcl"),
)
_SCNO = re.compile(r"^\s*\d+\s*-\s*.+?\s*-\s*elec\.?\s*([A-Za-z]\d+)\s*$", re.I)


def parse_sample_points_xlsx(content: bytes) -> list[dict[str, Any]]:
    """AriaLIMS sample-point list (SCID, SCNo, Location, UnitID, UnitTag) → point dicts.

    ``SCNo`` ("01-Brine- elec. A1") carries the electrolyzer, ``UnitTag`` the sample kind.
    """
    import io

    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = {str(v or "").strip().lower(): i for i, v in enumerate(rows[0])}
    for required in ("scid", "scno"):
        if required not in header:
            raise ValueError(f"Column '{required.upper()}' not found in the first row")
    tag_col = header.get("unittag")
    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        try:
            scid = int(row[header["scid"]])
        except (TypeError, ValueError):
            continue
        scno = str(row[header["scno"]] or "").strip()
        tag = str(row[tag_col] or "").strip().lower() if tag_col is not None else ""
        scno_l = scno.lower()
        analysis_type = next((t for needle, t in UNIT_TAG_TYPES if needle in tag), None) or next(
            (t for needle, t in UNIT_TAG_TYPES if needle in scno_l), None
        )
        match = _SCNO.match(scno)
        if not analysis_type or not match:
            out.append({"scid": scid, "name": scno, "error": "Could not read sample kind / electrolyzer from SCNo"})
            continue
        out.append(
            {
                "scid": scid,
                "name": scno,
                "analysis_type": analysis_type,
                "scope": "electrolyzer",
                "electrolyzer": match.group(1).upper(),
            }
        )
    return out


def import_sample_points(db: Session, entries: list[dict[str, Any]]) -> dict[str, Any]:
    """Create/update sampling points; keeps a point's own mapping and on/off switch."""
    created = updated = 0
    skipped: list[dict[str, Any]] = []
    for entry in entries:
        if entry.get("error"):
            skipped.append({"scid": entry["scid"], "name": entry.get("name"), "reason": entry["error"]})
            continue
        point = db.query(models.AriaLimsSamplingPoint).filter(models.AriaLimsSamplingPoint.scid == entry["scid"]).first()
        if point is None:
            db.add(models.AriaLimsSamplingPoint(parameter_map={}, enabled=True, **entry))
            created += 1
        else:
            for key in ("name", "analysis_type", "scope", "electrolyzer"):
                setattr(point, key, entry[key])
            updated += 1
    db.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "total": len(entries)}


def _norm_name(name: str) -> str:
    return re.sub(r"\s+", " ", str(name or "").strip()).lower()


def _plant_key(analysis_type: str, key: str) -> str:
    """The whole-plant forms drop the per-electrolyzer suffixes (" An", " Pb", " Cl")."""
    if analysis_type == "anolyte":
        if key == "density at 20° An":
            return "density at 20°"
        return key[:-3] if key.endswith(" An") else key
    if analysis_type == "pure_brine":
        return key[:-3] if key.endswith(" Pb") else key
    if analysis_type == "chlorine_gas":
        if key == "restgas Cl":
            return "Restgas"
        return key[:-3] if key.endswith(" Cl") else key
    return key


def mapped_key(analysis_type: str | None, scope: str | None, name: str, custom: dict[str, str]) -> str:
    """Parameter key an AriaLIMS analysis is stored under."""
    wanted = _norm_name(name)
    if wanted in custom:
        return custom[wanted]
    key = DEFAULT_PARAMETER_MAPS.get(analysis_type or "", {}).get(wanted)
    if not key:
        return str(name).strip()
    return _plant_key(analysis_type or "", key) if scope == "total_plant" else key


def default_settings() -> models.AriaLimsSyncSettings:
    return models.AriaLimsSyncSettings(
        enabled=False,
        base_url="",
        daily_time="01:00",
        lookback_days=7,
        analysis_types=",".join(ANALYSIS_TYPES.keys()),
    )


def get_or_create_settings(db: Session) -> models.AriaLimsSyncSettings:
    row = db.query(models.AriaLimsSyncSettings).first()
    if row:
        return row
    row = default_settings()
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def config_from_settings(row: models.AriaLimsSyncSettings) -> AriaLimsConfig:
    return AriaLimsConfig(
        base_url=normalize_base_url(row.base_url) or "",
        username=row.username,
        password=row.password,
        api_token=row.api_token,
    )


def selected_analysis_types(row: models.AriaLimsSyncSettings) -> list[str]:
    raw = (row.analysis_types or "").strip()
    if not raw:
        return list(ANALYSIS_TYPES.keys())
    wanted = [p.strip() for p in raw.replace(";", ",").split(",") if p.strip()]
    return [t for t in wanted if t in ANALYSIS_TYPES]


def date_window(*, end: date | None = None, lookback_days: int = 7, include_today: bool = True) -> tuple[date, date]:
    end_day = end or date.today()
    if not include_today:
        end_day = end_day - timedelta(days=1)
    days = max(1, int(lookback_days or 1))
    start_day = end_day - timedelta(days=days - 1)
    return start_day, end_day


def _value(raw: Any) -> Any:
    """'1.32' → 1.32; text such as '<0.1' stays text."""
    if isinstance(raw, (int, float)):
        return raw
    text = str(raw).strip() if raw is not None else ""
    if _NUMBER.match(text):
        try:
            return float(text.replace(",", "."))
        except ValueError:
            return text
    return text


def _sampled_at(raw: Any) -> datetime | None:
    text = str(raw or "").strip().replace("Z", "+00:00")
    if not text:
        return None
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError:
        try:
            parsed = datetime.fromisoformat(text[:10])
        except ValueError:
            return None
    return parsed.replace(tzinfo=None)


def point_label(rows: list[dict[str, Any]]) -> str | None:
    for row in rows:
        name = str(row.get("scno") or "").strip()
        if name:
            return name
    return None


def results_to_samples(
    point: models.AriaLimsSamplingPoint | None,
    rows: list[dict[str, Any]],
    *,
    analysis_type: str | None = None,
    window: tuple[date, date] | None = None,
) -> list[dict[str, Any]]:
    """Group AriaLIMS result rows into ``AnalysisSample`` field dicts (one per sampling time)."""
    mapping = {
        _norm_name(k): str(v).strip()
        for k, v in ((point.parameter_map if point else None) or {}).items()
        if str(v).strip()
    }
    the_type = (point.analysis_type if point else analysis_type) or ""
    the_scope = (point.scope if point else None) or "total_plant"
    by_time: dict[datetime, dict[str, Any]] = {}
    for row in rows:
        sampled = _sampled_at(row.get("samplingTime") or row.get("samplingtime"))
        name = str(row.get("analysisname") or "").strip()
        if sampled is None or not name:
            continue
        if window and not (window[0] <= sampled.date() <= window[1]):
            continue
        key = mapped_key(the_type, the_scope, name, mapping)
        by_time.setdefault(sampled, {})[key] = _value(row.get("value"))

    samples = []
    for sampled, params in sorted(by_time.items()):
        if the_type == "chlorine_gas":
            # The Access form has one "Cl2 + CO2" figure; AriaLIMS reports them separately.
            cl2, co2 = params.get("Cl2"), params.get("CO2")
            if isinstance(cl2, (int, float)) and isinstance(co2, (int, float)):
                params[_plant_key("chlorine_gas", CL2_CO2_KEY) if the_scope == "total_plant" else CL2_CO2_KEY] = round(
                    cl2 + co2, 4
                )
        samples.append(
            {
                "analysis_type": point.analysis_type if point else analysis_type,
                "scope": (point.scope if point else None) or "total_plant",
                "electrolyzer": point.electrolyzer if point else None,
                "position": point.position if point else None,
                "group_nr": point.group_nr if point else None,
                "sub_plant": point.sub_plant if point else None,
                "date": sampled.date().isoformat(),
                "time": sampled.strftime("%H:%M"),
                "parameters": params,
            }
        )
    return samples


def upsert_samples(db: Session, samples: list[dict[str, Any]]) -> dict[str, int]:
    """Insert samples; a sample for the same point and moment gets its parameters merged."""
    created = updated = 0
    for sample in samples:
        if not sample or not sample.get("analysis_type"):
            continue
        day = None
        if sample.get("date"):
            try:
                day = datetime.fromisoformat(str(sample["date"])[:10])
            except ValueError:
                day = None
        scope = sample.get("scope") or "total_plant"
        existing = (
            db.query(models.AnalysisSample)
            .filter(
                models.AnalysisSample.analysis_type == sample["analysis_type"],
                models.AnalysisSample.scope == scope,
                models.AnalysisSample.date == day,
                models.AnalysisSample.time == sample.get("time"),
                models.AnalysisSample.electrolyzer == sample.get("electrolyzer"),
                models.AnalysisSample.position == sample.get("position"),
                models.AnalysisSample.group_nr == sample.get("group_nr"),
                models.AnalysisSample.sub_plant == sample.get("sub_plant"),
            )
            .first()
        )
        incoming = sample.get("parameters") or {}
        if existing:
            merged = {**(existing.parameters or {}), **incoming}
            if merged != (existing.parameters or {}):
                existing.parameters = merged
                updated += 1
            continue
        db.add(
            models.AnalysisSample(
                analysis_type=sample["analysis_type"],
                scope=scope,
                electrolyzer=sample.get("electrolyzer"),
                position=sample.get("position"),
                group_nr=sample.get("group_nr"),
                sub_plant=sample.get("sub_plant"),
                date=day,
                time=sample.get("time"),
                parameters=incoming,
            )
        )
        created += 1
    db.commit()
    return {"created": created, "updated": updated}


def pull_from_arialims(
    db: Session,
    row: models.AriaLimsSyncSettings | None = None,
    *,
    include_today: bool = True,
    progress: ProgressCb = None,
) -> dict[str, Any]:
    """Pull every enabled sampling point. One failing point does not stop the others."""
    row = row or get_or_create_settings(db)
    cfg = config_from_settings(row)
    types = set(selected_analysis_types(row))
    start, end = date_window(lookback_days=row.lookback_days or 7, include_today=include_today)

    result: dict[str, Any] = {
        "ok": False,
        "mode": "api",
        "files_scanned": 0,
        "files_applied": 0,
        "rows_upserted": 0,
        "message": "",
        "details": [],
        "catalog": endpoint_catalog(),
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "contract_ready": True,
    }

    if not cfg.base_url:
        result["message"] = "AriaLIMS base URL is not configured (Settings → AriaLims Sync)."
        _mark_run(db, row, status="error", message=result["message"])
        return result

    points = [
        p
        for p in db.query(models.AriaLimsSamplingPoint).order_by(models.AriaLimsSamplingPoint.scid).all()
        if p.enabled and p.analysis_type in types
    ]
    if not points:
        result["message"] = "No sampling points configured. Add the SCIDs to pull under Settings → AriaLims Sync."
        _mark_run(db, row, status="error", message=result["message"])
        return result

    total = len(points)
    if progress:
        progress(0, total)
    failures = 0
    done = 0
    for start_idx in range(0, total, BATCH_SIZE):
        batch = points[start_idx : start_idx + BATCH_SIZE]
        try:
            # EndTime may be exclusive on the server: ask one day further, filter locally.
            rows = fetch_results(cfg, [p.scid for p in batch], start, end + timedelta(days=1))
        except AriaLimsNotConfigured as exc:
            result["message"] = str(exc)
            _mark_run(db, row, status="error", message=result["message"][:480])
            return result
        except Exception as exc:  # noqa: BLE001
            logger.warning("AriaLIMS batch of %s SCID(s) failed: %s", len(batch), exc)
            failures += len(batch)
            for point in batch:
                result["details"].append(
                    {"scid": point.scid, "name": point.name, "analysis_type": point.analysis_type, "status": "error", "error": str(exc)[:300]}
                )
            done += len(batch)
            if progress:
                progress(done, total)
            continue

        by_scid: dict[int, list[dict[str, Any]]] = {}
        for item in rows:
            try:
                by_scid.setdefault(int(item.get("scid")), []).append(item)
            except (TypeError, ValueError):
                continue
        for point in batch:
            detail: dict[str, Any] = {"scid": point.scid, "name": point.name, "analysis_type": point.analysis_type}
            point_rows = by_scid.get(point.scid, [])
            try:
                samples = results_to_samples(point, point_rows, window=(start, end))
                counts = upsert_samples(db, samples)
                result["rows_upserted"] += counts["created"] + counts["updated"]
                detail.update(
                    status="ok",
                    received=len(point_rows),
                    samples=len(samples),
                    created=counts["created"],
                    updated=counts["updated"],
                )
                if not point.name:
                    point.name = point_label(point_rows)
            except Exception as exc:  # noqa: BLE001
                logger.warning("AriaLIMS SCID %s failed: %s", point.scid, exc)
                db.rollback()
                failures += 1
                detail.update(status="error", error=str(exc)[:300])
            result["details"].append(detail)
            done += 1
            if progress:
                progress(done, total)

    db.commit()
    result["ok"] = failures == 0
    received = sum(int(d.get("received") or 0) for d in result["details"])
    result["message"] = (
        f"Updated {result['rows_upserted']} analysis sample(s) from {total - failures}/{total} sampling point(s) "
        f"({received} result row(s) received)."
        + ("" if not failures else f" {failures} point(s) failed: see details.")
    )
    _mark_run(db, row, status="success" if failures == 0 else "error", message=result["message"][:480])
    return result


def _mark_run(db: Session, row: models.AriaLimsSyncSettings, *, status: str, message: str) -> None:
    row.last_run_at = datetime.now(timezone.utc).replace(tzinfo=None)
    row.last_run_status = status
    row.last_run_message = message
    db.commit()


def ping(db: Session, row: models.AriaLimsSyncSettings) -> dict[str, Any]:
    probe = (
        db.query(models.AriaLimsSamplingPoint)
        .filter(models.AriaLimsSamplingPoint.enabled.is_(True))
        .order_by(models.AriaLimsSamplingPoint.scid)
        .first()
    )
    return test_connection(config_from_settings(row), probe.scid if probe else None)


def preview_point(db: Session, row: models.AriaLimsSyncSettings, scid: int, days: int = 7) -> dict[str, Any]:
    """Fetch one SCID without saving: what the API returns and how it would be grouped."""
    start, end = date_window(lookback_days=days, include_today=True)
    rows = fetch_results(config_from_settings(row), scid, start, end + timedelta(days=1))
    point = db.query(models.AriaLimsSamplingPoint).filter(models.AriaLimsSamplingPoint.scid == scid).first()
    in_window = [r for r in rows if (s := _sampled_at(r.get("samplingTime"))) and start <= s.date() <= end]
    units: dict[str, str] = {}
    for r in rows:
        name = str(r.get("analysisname") or "").strip()
        if name:
            units.setdefault(name, str(r.get("unitofmeaserment") or r.get("unitofmeasurement") or "").strip())
    return {
        "scid": scid,
        "name": point_label(rows),
        "date_from": start.isoformat(),
        "date_to": end.isoformat(),
        "received": len(rows),
        "analyses": [{"name": n, "unit": u} for n, u in units.items()],
        "rows": in_window[:60],
        "mapped": point is not None,
    }


def apply_manual_payload(db: Session, payload: dict | list, *, analysis_type: str) -> dict[str, Any]:
    """Ingest a pasted AriaLIMS response (``{"results": [...]}``). Known SCIDs use their mapping,
    others are stored as ``analysis_type`` for the whole plant."""
    if analysis_type not in ANALYSIS_TYPES:
        raise ValueError(f"Unsupported analysis_type '{analysis_type}'")
    rows = extract_results(payload)
    if not rows and not isinstance(payload, (dict, list)):
        raise ValueError("JSON must be an object or array")
    by_scid: dict[Any, list[dict[str, Any]]] = {}
    for item in rows:
        by_scid.setdefault(item.get("scid"), []).append(item)
    samples: list[dict[str, Any]] = []
    for scid, group in by_scid.items():
        point = None
        if scid is not None:
            point = db.query(models.AriaLimsSamplingPoint).filter(models.AriaLimsSamplingPoint.scid == int(scid)).first()
        samples.extend(results_to_samples(point, group, analysis_type=analysis_type))
    counts = upsert_samples(db, samples)
    return {"ok": True, "received": len(rows), "upserted": counts["created"] + counts["updated"], "analysis_type": analysis_type}
