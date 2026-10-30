"""Map an uploaded worksheet onto the form the user is currently in.

Headers may be field names, the English/Persian captions of that form, or the
same columns in form order with different titles. A sheet that matches nothing
is rejected instead of being written into the wrong fields.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
import uuid
import zipfile
from pathlib import Path
from datetime import date, datetime, time, timedelta
from types import UnionType
from typing import Any, Sequence, Union, get_args, get_origin

from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from pydantic import BaseModel, ValidationError
from sqlalchemy import Integer
from sqlalchemy.orm import Session

from .database import commit_db

_DROP = re.compile(r"[\s_\-./\\()\[\]{}%#:;,'\"`+]+")
_JALALI = re.compile(
    r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?\s*(AM|PM)?)?$",
    re.IGNORECASE,
)

# Captions that do not compact to the stored field name.
_ALIASES: dict[str, str] = {
    "findings": "finding",
    "befund": "finding",
    "despatch": "dispatch_date",
    "dispatch": "dispatch_date",
    "shippingdate": "dispatch_date",
    "shipping": "dispatch_date",
    "versand": "dispatch_date",
    "return": "return_date",
    "returndate": "return_date",
    "sheetreturn": "return_date",
    "ruckgabe": "return_date",
    "rueckgabe": "return_date",
    "element": "element_nr",
    "elementno": "element_nr",
    "cell": "cell",
    "stage": "stage",
    "group": "group_nr",
    "groupno": "group_nr",
    "membrane": "membrane_nr",
    "anode": "anode_nr",
    "anodeno": "anode_nr",
    "anodenumber": "anode_nr",
    "anodenr": "anode_nr",
    "anodennummer": "anode_nr",
    "cathode": "cathode_nr",
    "kathode": "cathode_nr",
    "cathodeno": "cathode_nr",
    "cathodenumber": "cathode_nr",
    "cathodenr": "cathode_nr",
    "kathodennummer": "cathode_nr",
    "membraneno": "membrane_nr",
    "membranenumber": "membrane_nr",
    "membranenr": "membrane_nr",
    "membranetype": "membrane_type",
    "coatingthickness": "residual_thickness",
    "coatingthicknesspct": "residual_thickness",
    "thickness": "residual_thickness",
    "potentialv": "potential",
    "dateofinspection": "check_date",
    "inspectiondate": "check_date",
    "partofplant": "plant_part",
    "plantpart": "plant_part",
    "anlagenteil": "plant_part",
    "reason": "cause",
    "ursache": "cause",
    "remark": "remarks",
    "comment": "remarks",
    "comments": "remarks",
    "bemerkung": "remarks",
    "no": "nr",
    "number": "nr",
    "repair": "repair_work",
    "repairwork": "repair_work",
    "hersteller": "manufacturer",
    "aktion": "action",
    "شمارهآند": "anode_nr",
    "شمارهکاتد": "cathode_nr",
    "شمارهغشا": "membrane_nr",
    "یافتهها": "finding",
    "ارسال": "dispatch_date",
    "تاریخارسال": "dispatch_date",
    "بازگشت": "return_date",
    "تاریخبازگشت": "return_date",
    "ضخامتپوشش": "residual_thickness",
    "تاریخبازرسی": "check_date",
    "تاریخدریافت": "received_date",
    "بخشواحد": "plant_part",
    "زمانتوقف": "shutdown_time",
    "زمانراهاندازی": "startup_time",
    "علت": "cause",
    "دسته": "category",
    "دستهبندی": "category",
    "ملاحظات": "remarks",
    "کد": "code",
    "شماره": "nr",
    "تاریخ": "date",
}


def compact(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value)).replace("\u200c", "")
    text = "".join(ch for ch in text.lower() if ch.isalnum())
    return _DROP.sub("", text)


def _keys(label: str) -> set[str]:
    base = compact(label)
    if not base:
        return set()
    keys = {base}
    if "number" in base:
        keys.add(base.replace("number", "nr"))
    if base.endswith("no") and len(base) > 2:
        keys.add(base[:-2] + "nr")
    if base.endswith("s") and len(base) > 4:
        keys.add(base[:-1])
    elif len(base) > 3:
        keys.add(base + "s")
    return keys


def _field_index(fields: Sequence[str]) -> dict[str, str]:
    index: dict[str, str] = {}
    for field in fields:
        for key in _keys(field.replace("_", " ")):
            index.setdefault(key, field)
        index.setdefault(compact(field), field)
    for key, target in _ALIASES.items():
        if target in fields:
            index.setdefault(key, target)
    return index


def map_headers(
    header: Sequence[Any],
    fields: Sequence[str],
    optional: Sequence[str] = (),
) -> tuple[dict[int, str], bool, list[str]]:
    """Return column-index to field, whether columns were placed by form order, and ignored titles."""
    form_fields = list(fields)
    allowed = list(dict.fromkeys([*form_fields, *optional]))
    index = _field_index(allowed)
    filled = [(i, cell) for i, cell in enumerate(header) if cell is not None and str(cell).strip()]
    if not filled:
        raise HTTPException(status_code=400, detail="The Excel sheet has no column titles.")

    mapping: dict[int, str] = {}
    used: set[str] = set()
    ignored: list[str] = []
    date_fields = [f for f in allowed if f == "date" or f.endswith("_date") or f.endswith("_time")]
    for i, cell in filled:
        keys = _keys(str(cell))
        field = next((index[key] for key in keys if key in index and index[key] not in used), None)
        if field is None and keys & {"date", "تاریخ"} and "date" not in used:
            if "date" in allowed:
                field = "date"
            elif len(date_fields) == 1:
                field = date_fields[0]
        if field is None or field in used:
            ignored.append(str(cell).strip())
            continue
        mapping[i] = field
        used.add(field)

    if mapping:
        return mapping, False, ignored

    if len(filled) == len(form_fields):
        return {i: field for (i, _), field in zip(filled, form_fields)}, True, []

    names = ", ".join(form_fields)
    raise HTTPException(
        status_code=400,
        detail=(
            "Excel columns do not match this form. "
            f"Expected columns: {names}. "
            "سرستون‌های فایل با فیلدهای این قسمت جور نیست. "
            "نام ستون‌ها را مطابق همین فرم بگذارید، یا ستون‌ها را به همان تعداد و به ترتیب فیلدهای فرم بچینید."
        ),
    )


def _load_rows(content: bytes) -> list[tuple]:
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail="The file is not a readable Excel workbook.") from exc
    try:
        ws = wb.active
        return [tuple(row) for row in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def parse_sheet(
    content: bytes,
    fields: Sequence[str],
    optional: Sequence[str] = (),
) -> tuple[dict[int, str], list[tuple], bool, list[str]]:
    rows = _load_rows(content)
    if not rows:
        raise HTTPException(status_code=400, detail="The Excel sheet is empty.")

    index = _field_index(list(dict.fromkeys([*fields, *optional])))
    best_at = 0
    best_score = -1
    scan = rows[:8]
    for offset, row in enumerate(scan):
        score = 0
        used: set[str] = set()
        for cell in row:
            if cell is None or not str(cell).strip():
                continue
            field = next((index[key] for key in _keys(str(cell)) if key in index and index[key] not in used), None)
            if field:
                used.add(field)
                score += 1
        if score > best_score:
            best_score = score
            best_at = offset
    if best_score <= 0:
        for offset, row in enumerate(scan):
            filled = sum(1 for cell in row if cell is not None and str(cell).strip())
            if filled >= 2:
                best_at = offset
                break

    mapping, arranged, ignored = map_headers(rows[best_at], fields, optional)
    data = [row for row in rows[best_at + 1 :] if any(cell is not None and str(cell).strip() != "" for cell in row)]
    return mapping, data, arranged, ignored


def _unwrap(annotation: Any) -> Any:
    origin = get_origin(annotation)
    if origin in (UnionType, Union):
        args = [arg for arg in get_args(annotation) if arg is not type(None)]
        return args[0] if args else str
    return annotation


def jalali_to_gregorian(jy: int, jm: int, jd: int) -> tuple[int, int, int]:
    jy += 1595
    days = -355668 + (365 * jy) + ((jy // 33) * 8) + (((jy % 33) + 3) // 4) + jd
    if jm < 7:
        days += (jm - 1) * 31
    else:
        days += ((jm - 7) * 30) + 186
    gy = 400 * (days // 146097)
    days %= 146097
    if days > 36524:
        gy += 100 * ((days - 1) // 36524)
        days = (days - 1) % 36524
        if days >= 365:
            days += 1
    gy += 4 * (days // 1461)
    days %= 1461
    if days > 365:
        gy += (days - 1) // 365
        days = (days - 1) % 365
    gd = days + 1
    leap = gy % 4 == 0 and (gy % 100 != 0 or gy % 400 == 0)
    month_days = [0, 31, 29 if leap else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gm = 1
    while gm <= 12 and gd > month_days[gm]:
        gd -= month_days[gm]
        gm += 1
    return gy, gm, gd


def _as_datetime(value: Any) -> datetime | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return datetime(value.year, value.month, value.day)
    if isinstance(value, time):
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        serial = float(value)
        if 1 <= serial < 80000:
            return datetime(1899, 12, 30) + timedelta(days=serial)
        return None
    text = str(value).strip()
    if not text:
        return None
    match = _JALALI.match(text)
    if match:
        year, month, day = int(match.group(1)), int(match.group(2)), int(match.group(3))
        hour = int(match.group(4) or 0)
        minute = int(match.group(5) or 0)
        second = int(match.group(6) or 0)
        ampm = (match.group(7) or "").upper()
        if ampm == "PM" and hour < 12:
            hour += 12
        if ampm == "AM" and hour == 12:
            hour = 0
        if 1200 <= year <= 1600:
            year, month, day = jalali_to_gregorian(year, month, day)
        return datetime(year, month, day, hour, minute, second)
    try:
        return datetime.fromisoformat(text.replace("Z", "+00:00").replace("/", "-")).replace(tzinfo=None)
    except ValueError:
        return None


def _coerce(value: Any, annotation: Any) -> Any:
    if value is None or value == "":
        return None
    target = _unwrap(annotation)
    if target in (datetime, date):
        try:
            parsed = _as_datetime(value)
        except (ValueError, TypeError, OverflowError):
            return None
        if target is date and parsed is not None:
            return parsed.date()
        return parsed
    if target is int:
        if isinstance(value, bool):
            return int(value)
        if isinstance(value, float):
            return int(value)
        return int(str(value).strip())
    if target is float:
        if isinstance(value, str):
            return float(value.strip().replace(",", "."))
        return float(value)
    if target is bool:
        if isinstance(value, str):
            return value.strip().lower() in {"1", "true", "yes", "y", "بله"}
        return bool(value)
    if target is str:
        if isinstance(value, float) and value.is_integer():
            return str(int(value))
        if isinstance(value, datetime):
            return value.isoformat(sep=" ")
        return str(value).strip()
    return value


def _pk_value(model: type, pk_field: str, value: Any) -> Any:
    if value is None or value == "":
        return None
    try:
        column = model.__table__.columns[pk_field]
        if isinstance(column.type, Integer) and not isinstance(value, bool):
            return int(float(value))
    except Exception:  # noqa: BLE001
        pass
    if isinstance(value, float) and value.is_integer():
        return str(int(value)) if not isinstance(value, int) else value
    if isinstance(value, str):
        return value.strip()
    return value


async def read_xlsx(file: UploadFile) -> bytes:
    """Keep the old helper: some callers still want the bytes on the request thread."""
    path, _name = await stash_upload(file)
    try:
        return path.read_bytes()
    finally:
        path.unlink(missing_ok=True)


def _uploads_dir() -> Path:
    from .paths import data_root

    folder = data_root() / "uploads_tmp"
    folder.mkdir(parents=True, exist_ok=True)
    return folder


def is_elements_workbook_name(filename: str | None) -> bool:
    """Program Assembly Data export is named elements.xlsx (Excel may add ' (3)')."""
    stem = Path(filename or "").name.lower()
    stem = Path(stem).stem
    compact_name = re.sub(r"[^a-z0-9]+", "", stem)
    return compact_name.startswith("elements") or "elements" in compact_name


def _assert_xlsx(path: Path, size: int) -> None:
    if size < 32:
        path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail="The uploaded file is empty or was cut off.")
    if not zipfile.is_zipfile(path):
        path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=400,
            detail="The Excel file arrived incomplete (not a valid .xlsx archive). Upload it again.",
        )
    with zipfile.ZipFile(path) as book:
        bad = book.testzip()
        if bad:
            path.unlink(missing_ok=True)
            raise HTTPException(
                status_code=400,
                detail=f"The Excel file arrived incomplete (broken part: {bad}). Upload it again.",
            )


def begin_chunked_upload(filename: str, size: int, chunk_size: int | None = None) -> dict[str, Any]:
    name = (filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File must be an .xlsx Excel workbook.")
    upload_id = uuid.uuid4().hex
    folder = _uploads_dir()
    dest = folder / f"{upload_id}.part"
    dest.write_bytes(b"")
    meta = {
        "filename": Path(filename).name,
        "size": int(size),
        "received": 0,
        "chunk_size": int(chunk_size or 512 * 1024),
    }
    (folder / f"{upload_id}.json").write_text(json.dumps(meta), encoding="utf-8")
    return {"upload_id": upload_id, "chunk_size": meta["chunk_size"]}


def write_chunk(upload_id: str, offset: int, data: bytes) -> dict[str, Any]:
    folder = _uploads_dir()
    meta_path = folder / f"{upload_id}.json"
    dest = folder / f"{upload_id}.part"
    if not meta_path.is_file() or not dest.is_file():
        raise HTTPException(status_code=404, detail="Upload session not found.")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    expected = int(meta["size"])
    if offset < 0 or offset > expected or offset + len(data) > expected + 64:
        raise HTTPException(status_code=400, detail="Upload chunk is outside the declared file size.")
    with dest.open("r+b") as out:
        out.seek(offset)
        out.write(data)
        received = out.tell()
        if dest.stat().st_size < received:
            pass
    meta["received"] = max(int(meta.get("received") or 0), offset + len(data))
    meta_path.write_text(json.dumps(meta), encoding="utf-8")
    return {"received": meta["received"], "size": expected}


def take_chunked_upload(upload_id: str) -> tuple[Path, str]:
    folder = _uploads_dir()
    meta_path = folder / f"{upload_id}.json"
    dest = folder / f"{upload_id}.part"
    if not meta_path.is_file() or not dest.is_file():
        raise HTTPException(status_code=404, detail="Upload session not found.")
    meta = json.loads(meta_path.read_text(encoding="utf-8"))
    size = dest.stat().st_size
    expected = int(meta.get("size") or 0)
    if expected and size != expected:
        dest.unlink(missing_ok=True)
        meta_path.unlink(missing_ok=True)
        raise HTTPException(
            status_code=400,
            detail=f"The Excel file arrived incomplete ({size} of {expected} bytes). Upload it again.",
        )
    _assert_xlsx(dest, size)
    filename = meta.get("filename") or "upload.xlsx"
    final = folder / f"{upload_id}_{Path(filename).name}"
    dest.replace(final)
    meta_path.unlink(missing_ok=True)
    return final, filename


async def resolve_upload(file: UploadFile | None, upload_id: str | None) -> tuple[Path, str]:
    if upload_id:
        return take_chunked_upload(upload_id)
    if file is None:
        raise HTTPException(status_code=400, detail="No Excel file was sent.")
    return await stash_upload(file)


async def stash_upload(file: UploadFile):
    """Write the upload to disk in 512 KB chunks so a large Assembly workbook is not cut off in memory."""
    name = (file.filename or "").lower()
    if not name.endswith((".xlsx", ".xlsm")):
        raise HTTPException(status_code=400, detail="File must be an .xlsx Excel workbook.")
    dest = _uploads_dir() / f"{uuid.uuid4().hex}_{Path(file.filename or 'upload.xlsx').name}"
    size = 0
    try:
        with dest.open("wb") as out:
            while True:
                chunk = await file.read(512 * 1024)
                if not chunk:
                    break
                out.write(chunk)
                size += len(chunk)
    except Exception:
        dest.unlink(missing_ok=True)
        raise
    _assert_xlsx(dest, size)
    return dest, file.filename or dest.name


def import_excel_bytes(
    db: Session,
    model: type,
    write_schema: type[BaseModel],
    content: bytes,
    *,
    pk_field: str = "id",
    mode: str = "merge",
    prepare: Any = None,
    progress: Any = None,
    file_key: str | None = None,
    file_name: str | None = None,
    uploaded_by: str | None = None,
) -> dict[str, Any]:
    write_fields = list(write_schema.model_fields.keys())
    optional = [pk_field] if pk_field not in write_fields else []
    mapping, data_rows, arranged, ignored = parse_sheet(content, write_fields, optional)
    if mode == "replace":
        db.query(model).delete()
        commit_db(db)

    created = updated = skipped = 0
    errors: list[str] = []
    total = len(data_rows)
    if progress:
        progress(0, total)
    for done, row in enumerate(data_rows, start=1):
        raw: dict[str, Any] = {}
        for column, field in mapping.items():
            if column >= len(row):
                continue
            value = row[column]
            if value is None or value == "":
                continue
            raw[field] = value
        if not raw:
            continue
        pk_raw = raw.get(pk_field)
        payload_input: dict[str, Any] = {}
        for key, value in raw.items():
            if key not in write_schema.model_fields:
                continue
            try:
                payload_input[key] = _coerce(value, write_schema.model_fields[key].annotation)
            except (ValidationError, ValueError, TypeError) as exc:
                if len(errors) < 8:
                    errors.append(f"row {done} {key}: {str(exc).split(chr(10))[0][:200]}")
                continue
        if not payload_input and pk_raw is None:
            continue
        try:
            payload = write_schema.model_validate(payload_input)
            data = payload.model_dump(exclude_unset=True)
            pk_value = _pk_value(model, pk_field, data.get(pk_field, pk_raw))
        except (ValidationError, ValueError, TypeError) as exc:
            skipped += 1
            if len(errors) < 8:
                errors.append(str(exc).split("\n")[0][:300])
            continue

        existing = None
        if pk_value is not None:
            existing = db.query(model).filter(getattr(model, pk_field) == pk_value).first()
        if existing is not None:
            for key, value in data.items():
                if key == pk_field:
                    continue
                setattr(existing, key, value)
            if prepare:
                prepare(existing)
            updated += 1
        else:
            obj = model(**{k: v for k, v in data.items() if k != pk_field or pk_field in write_fields})
            if pk_value is not None and pk_field not in data:
                setattr(obj, pk_field, pk_value)
            if prepare:
                prepare(obj)
            db.add(obj)
            created += 1
        if (created + updated) % 200 == 0:
            commit_db(db)
        if progress and (done == total or done % 25 == 0):
            progress(done, total)

    if created + updated == 0:
        db.rollback()
        detail = "No rows were imported."
        if errors:
            detail = f"{detail} {errors[0]}"
        raise HTTPException(status_code=400, detail=detail)

    commit_db(db)
    if file_key:
        from .table_files import remember_workbook

        remember_workbook(db, file_key, file_name, content, db.query(model).count(), uploaded_by)
    return {
        "ok": True,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "arranged": arranged,
        "mapped": [mapping[i] for i in sorted(mapping)],
        "ignored": ignored,
        "errors": errors,
    }


def _self_check() -> None:
    fields = ["anode_nr", "date", "finding", "action", "dispatch_date", "return_date"]
    mapping, arranged, ignored = map_headers(
        ["Anode No.", "Date", "Findings", "Action", "Despatch", "Return"],
        fields,
    )
    assert [mapping[i] for i in sorted(mapping)] == fields, mapping
    assert arranged is False and ignored == []
    mapping, arranged, _ = map_headers(["A", "B", "C", "D", "E", "F"], fields)
    assert arranged is True and [mapping[i] for i in range(6)] == fields
    try:
        map_headers(["Only one"], fields)
        raise AssertionError("mismatched sheet should fail")
    except HTTPException as exc:
        assert exc.status_code == 400
    assert jalali_to_gregorian(1405, 7, 6) == (2026, 9, 28)
    assert jalali_to_gregorian(1404, 1, 1) == (2025, 3, 21)
    shutdown = ["nr", "plant_part", "shutdown_time", "startup_time", "code", "cause", "category", "remarks"]
    mapping, _, ignored = map_headers(
        ["No.", "Part of Plant", "Shut Down Time", "Start Up Time", "Duration [h]", "Code", "Reason", "Category", "Remarks"],
        shutdown,
    )
    assert mapping[0] == "nr"
    assert mapping[1] == "plant_part"
    assert mapping[2] == "shutdown_time"
    assert mapping[6] == "cause"
    assert "Duration [h]" in ignored
