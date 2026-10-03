"""Shared helpers to export a list of ORM/dict rows to Excel (.xlsx) or PDF."""
import io
from datetime import date, datetime
from pathlib import Path
from typing import Any

from fastapi import Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

_PDF_FONT = "Helvetica"
_PDF_FONT_BOLD = "Helvetica-Bold"


def _ensure_pdf_fonts() -> tuple[str, str]:
    """Prefer Calibri (Windows) for clearer PDF table text."""
    global _PDF_FONT, _PDF_FONT_BOLD
    if _PDF_FONT == "Calibri":
        return _PDF_FONT, _PDF_FONT_BOLD
    candidates = [
        Path(r"C:\Windows\Fonts\calibri.ttf"),
        Path("/usr/share/fonts/truetype/crosextra/Carlito-Regular.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf"),
    ]
    bold_candidates = [
        Path(r"C:\Windows\Fonts\calibrib.ttf"),
        Path("/usr/share/fonts/truetype/crosextra/Carlito-Bold.ttf"),
        Path("/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"),
    ]
    regular = next((p for p in candidates if p.exists()), None)
    bold = next((p for p in bold_candidates if p.exists()), None)
    if regular is not None:
        try:
            pdfmetrics.registerFont(TTFont("Calibri", str(regular)))
            _PDF_FONT = "Calibri"
            if bold is not None:
                pdfmetrics.registerFont(TTFont("Calibri-Bold", str(bold)))
                _PDF_FONT_BOLD = "Calibri-Bold"
            else:
                _PDF_FONT_BOLD = "Calibri"
        except Exception:
            pass
    return _PDF_FONT, _PDF_FONT_BOLD


def _cell_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, (dict, list)):
        return str(value)
    return value


def infer_date_fields(fields: list[str]) -> list[str]:
    out: list[str] = []
    for f in fields:
        low = f.lower()
        if low.endswith("_date") or low.endswith("_time") or low in ("date", "timestamp", "created_at", "updated_at"):
            out.append(f)
    return out


def build_export_meta(all_fields: list[str]) -> dict[str, list[str]]:
    return {"fields": all_fields, "date_fields": infer_date_fields(all_fields)}


def pick_export_fields(all_fields: list[str], columns: str | None) -> list[str]:
    if not columns or not str(columns).strip():
        return all_fields
    requested = [c.strip() for c in str(columns).split(",") if c.strip()]
    picked = [f for f in requested if f in all_fields]
    return picked or all_fields


def _parse_row_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    if not text:
        return None
    if "T" in text:
        text = text.split("T", 1)[0]
    if " " in text and len(text) > 10:
        text = text.split(" ", 1)[0]
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def filter_rows_by_date(
    rows: list[dict[str, Any]],
    date_field: str | None,
    date_from: date | None,
    date_to: date | None,
) -> list[dict[str, Any]]:
    if not date_field or (date_from is None and date_to is None):
        return rows
    if date_field not in (rows[0].keys() if rows else []):
        return rows
    out: list[dict[str, Any]] = []
    for row in rows:
        row_date = _parse_row_date(row.get(date_field))
        if row_date is None:
            continue
        if date_from is not None and row_date < date_from:
            continue
        if date_to is not None and row_date > date_to:
            continue
        out.append(row)
    return out


def apply_export_filters(
    rows: list[dict[str, Any]],
    all_fields: list[str],
    *,
    columns: str | None = None,
    date_field: str | None = None,
    date_from: date | None = None,
    date_to: date | None = None,
) -> tuple[list[dict[str, Any]], list[str]]:
    fields = pick_export_fields(all_fields, columns)
    filtered = filter_rows_by_date(rows, date_field, date_from, date_to)
    return filtered, fields


class ExportFilters:
    """Shared query params for Excel/PDF export endpoints."""

    def __init__(
        self,
        columns: str | None = Query(default=None, description="Comma-separated export column keys"),
        date_field: str | None = Query(default=None),
        date_from: date | None = Query(default=None),
        date_to: date | None = Query(default=None),
    ):
        self.columns = columns
        self.date_field = date_field
        self.date_from = date_from
        self.date_to = date_to

    def apply(self, rows: list[dict[str, Any]], all_fields: list[str]) -> tuple[list[dict[str, Any]], list[str]]:
        return apply_export_filters(
            rows,
            all_fields,
            columns=self.columns,
            date_field=self.date_field,
            date_from=self.date_from,
            date_to=self.date_to,
        )


def rows_to_dicts(items: list, fields: list[str]) -> list[dict[str, Any]]:
    out = []
    for item in items:
        row = {}
        for f in fields:
            row[f] = _cell_value(getattr(item, f, None))
        out.append(row)
    return out


def export_xlsx(rows: list[dict[str, Any]], fields: list[str], title: str) -> StreamingResponse:
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Data"
    header_fill = PatternFill(start_color="0A246A", end_color="0A246A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    ws.append(fields)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in rows:
        ws.append([row.get(f) for f in fields])
    for i, f in enumerate(fields, start=1):
        max_len = max([len(str(f))] + [len(str(r.get(f, ""))) for r in rows[:200]])
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(max(max_len + 2, 10), 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"{title or 'export'}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


def export_pdf(rows: list[dict[str, Any]], fields: list[str], title: str) -> StreamingResponse:
    font, font_bold = _ensure_pdf_fonts()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4),
        leftMargin=1 * cm,
        rightMargin=1 * cm,
        topMargin=1 * cm,
        bottomMargin=1 * cm,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ExportTitle",
        parent=styles["Title"],
        fontName=font_bold,
        fontSize=14,
        leading=18,
        textColor=colors.HexColor("#0A246A"),
        spaceAfter=6,
    )
    cell_style = ParagraphStyle(
        "ExportCell",
        fontName=font,
        fontSize=8,
        leading=10,
        textColor=colors.black,
    )
    header_style = ParagraphStyle(
        "ExportHeader",
        fontName=font_bold,
        fontSize=8,
        leading=10,
        textColor=colors.white,
    )
    elements = [Paragraph(title, title_style), Spacer(1, 0.35 * cm)]

    def _p(text: str, style: ParagraphStyle) -> Paragraph:
        safe = str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    data = [[_p(f, header_style) for f in fields]]
    for row in rows:
        data.append([_p("" if row.get(f) is None else row.get(f), cell_style) for f in fields])
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0A246A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), font_bold),
                ("FONTNAME", (0, 1), (-1, -1), font),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#808080")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6FB")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    buf.seek(0)
    filename = f"{title or 'export'}.pdf"
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )
