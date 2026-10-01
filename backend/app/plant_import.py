"""Parsers for plant Excel files used at PVC Arvand (SiteMan voltage, LIMS lab, assembly, TAFKIK)."""
from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

# ---------------------------------------------------------------------------
# Jalali (Persian) calendar helpers — sufficient for plant dates 1300–1500
# ---------------------------------------------------------------------------

def _div(a: int, b: int) -> int:
    return a // b


def jalali_to_gregorian(jy: int, jm: int, jd: int) -> date:
    """Convert Jalali y/m/d to Gregorian date (algorithm from jalaali-js)."""
    if jy > 979:
        gy = 1600
        jy -= 979
    else:
        gy = 621
    days = 365 * jy + _div(jy, 33) * 8 + _div((jy % 33) + 3, 4) + 78 + jd + (31 * (jm - 1) if jm < 7 else (186 + 30 * (jm - 7)))
    gy += 400 * _div(days, 146097)
    days %= 146097
    if days > 36524:
        gy += 100 * _div(days - 1, 36524)
        days = (days - 1) % 36524
        if days >= 365:
            days += 1
    gy += 4 * _div(days, 1461)
    days %= 1461
    if days > 365:
        gy += _div(days - 1, 365)
        days = (days - 1) % 365
    gd = days + 1
    sal_a = [0, 31, 29 if (gy % 4 == 0 and gy % 100 != 0) or (gy % 400 == 0) else 28, 31, 30, 31, 30, 31, 31, 30, 31, 30, 31]
    gm = 1
    while gm <= 12 and gd > sal_a[gm]:
        gd -= sal_a[gm]
        gm += 1
    return date(gy, gm, gd)


def parse_plant_date(value: Any) -> date | None:
    """Parse Jalali or Gregorian dates from plant Excel cells."""
    if value is None or value == "" or value == "*":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = re.sub(r"\.{2,}", ".", str(value).strip().strip("."))
    if not text or text in {"*", "―", "-", "—"}:
        return None
    # Excel serial number
    if isinstance(value, (int, float)) and 20000 < float(value) < 80000:
        try:
            return (datetime(1899, 12, 30) + timedelta(days=float(value))).date()
        except Exception:
            pass
    # YYYY-MM-DD or YYYY/MM/DD (Gregorian if year > 1600)
    m = re.match(r"^(\d{4})[./\-](\d{1,2})[./\-](\d{1,2})", text)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        if y >= 1600:
            try:
                return date(y, mo, d)
            except ValueError:
                return None
        if 1300 <= y <= 1500:
            try:
                return jalali_to_gregorian(y, mo, d)
            except Exception:
                return None
    # YY.MM.DD Jalali short (e.g. 00.11.18 → 1400/11/18 if yy<=80 else 13yy)
    m = re.match(r"^(\d{1,2})[./\-](\d{1,2})[./\-](\d{1,2})$", text)
    if m:
        yy, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y = 1400 + yy if yy <= 80 else 1300 + yy
        try:
            return jalali_to_gregorian(y, mo, d)
        except Exception:
            return None
    # Persian month name + year e.g. آبان 1397
    persian_months = {
        "فروردین": 1, "فروزدین": 1,
        "اردیبهشت": 2, "اردیبهشتا": 2,
        "خرداد": 3,
        "تیر": 4,
        "مرداد": 5,
        "شهریور": 6,
        "مهر": 7,
        "آبان": 8, "ابان": 8,
        "آذر": 9, "اذر": 9,
        "دی": 10,
        "بهمن": 11,
        "اسفند": 12,
    }
    for name, mo in persian_months.items():
        if name in text:
            ym = re.search(r"(13|14)\d{2}", text)
            if not ym:
                continue
            # ARIAORMS header: "یکشنبه 5 ام مهر ماه 1405"
            day = 1
            day_m = re.search(rf"(\d{{1,2}})\s*(?:ام|م)?\s*{re.escape(name)}", text)
            if day_m:
                day = int(day_m.group(1))
            else:
                day_m = re.search(r"(\d{1,2})\s*ام", text)
                if day_m:
                    day = int(day_m.group(1))
            try:
                return jalali_to_gregorian(int(ym.group(0)), mo, day)
            except Exception:
                return None
    return None


def parse_plant_datetime(text: Any) -> datetime | None:
    if text is None:
        return None
    if isinstance(text, datetime):
        return text
    s = str(text).strip()
    # 1405/06/03 08:00
    m = re.match(r"^(\d{4})[./\-](\d{1,2})[./\-](\d{1,2})(?:\s+(\d{1,2}):(\d{2}))?", s)
    if m:
        d = parse_plant_date(f"{m.group(1)}/{m.group(2)}/{m.group(3)}")
        if not d:
            return None
        hh = int(m.group(4) or 0)
        mm = int(m.group(5) or 0)
        return datetime(d.year, d.month, d.day, hh, mm)
    d = parse_plant_date(s)
    if d:
        return datetime(d.year, d.month, d.day)
    return None


def parse_float_plant(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    text = str(value).strip()
    # SiteMan: "3/02\nبهزاد..." or "494/4 بهزاد"
    text = text.split("\n", 1)[0].strip()
    text = re.split(r"\s+", text, maxsplit=1)[0]
    text = text.replace(",", ".").replace("/", ".")
    text = re.sub(r"[^\d.\-]", "", text)
    if not text or text in {".", "-", "-."}:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _sheet_rows(wb, sheet_name: str | None = None) -> list[list[Any]]:
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ---------------------------------------------------------------------------
# 1) SiteMan / F2 / ARIAORMS LogSheets voltage Excel
# ---------------------------------------------------------------------------

_CELL_RE = re.compile(r"^(?P<el>[A-Za-z]+\d*)-(?P<pos>\d{1,3})\b")
_RECT_ITEM_RE = re.compile(r"^(?P<el>[A-Za-z]+\d+)\s*\[(?P<what>[^\]]+)\]", re.I)
_DATE_ROW_RE = re.compile(r"^\d{4}[./\-]\d{1,2}[./\-]\d{1,2}")
_EL_RE = re.compile(r"^(?P<el>[A-Za-z]+\d*)\[", re.I)
_EL_VOLTAGE_RE = re.compile(r"electrolyzer\s+voltage\s+(?P<el>[A-Za-z]+\d*)", re.I)
_EL_IN_TEXT_RE = re.compile(r"(?:ELECTROLYZER|EL)\s*(?P<el>[A-Za-z]+\d*)", re.I)


def _col_index(vals_lower: list[str], *names: str) -> int | None:
    for name in names:
        try:
            return vals_lower.index(name.lower())
        except ValueError:
            continue
    return None


def _detect_electrolyzer(*texts: str) -> str | None:
    for text in texts:
        if not text:
            continue
        for rx in (_EL_VOLTAGE_RE, _EL_RE, _EL_IN_TEXT_RE, _CELL_RE):
            m = rx.search(text)
            if m:
                return m.group("el").upper()
    return None


def parse_voltage_excel(content: bytes) -> dict[str, Any]:
    """Parse SiteMan / F2 / ARIAORMS LogSheetsReports Excel exports.

    ARIAORMS wide layout (typical):
      No | Unit | Equipment | Tag | Parameter | Unit | Normal Range | 02:00 | 06:00 | …
    Rows: "Electrolyzer voltage A2" (total) + "A2-001" … element cells.
    """
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    # Prefer clean numeric sheet over SiteMan annotated sheet
    preferred = None
    for name in wb.sheetnames:
        lower = name.lower()
        if "siteman" in lower or "whit" in lower:
            continue
        preferred = name
        break
    if preferred is None:
        preferred = wb.sheetnames[0]

    rows = _sheet_rows(wb, preferred)
    wb.close()

    reading_date: date | None = None
    for row in rows[:12]:
        for cell in row or []:
            if cell is None:
                continue
            d = parse_plant_date(cell)
            if d:
                reading_date = d
                break
        if reading_date:
            break

    header_idx = None
    time_cols: dict[int, str] = {}
    tag_col: int | None = None
    param_col: int | None = None
    equip_col: int | None = None
    for i, row in enumerate(rows):
        vals = [str(c).strip() if c is not None else "" for c in row]
        vals_lower = [v.lower() for v in vals]
        if "parameter" not in vals_lower and "tag" not in vals_lower:
            continue
        tag_col = _col_index(vals_lower, "tag")
        param_col = _col_index(vals_lower, "parameter")
        equip_col = _col_index(vals_lower, "equipment")
        # SiteMan legacy: Parameter often in column 1 without named Tag
        if param_col is None and tag_col is None and len(vals) > 1:
            param_col = 1
            equip_col = equip_col if equip_col is not None else 0
        for j, v in enumerate(vals):
            if re.match(r"^\d{1,2}:\d{2}$", v):
                time_cols[j] = v
        if time_cols and (param_col is not None or tag_col is not None):
            header_idx = i
            break

    if header_idx is None:
        return {"error": "Header row with Parameter/Tag / time columns not found", "readings": [], "totals": {}}

    electrolyzer = None
    readings: list[dict[str, Any]] = []
    totals: dict[str, Any] = {}
    total_slots: dict[tuple[str, str], dict[str, Any]] = {}
    operators: set[str] = set()
    out_of_range: list[dict[str, Any]] = []
    _date_row = _DATE_ROW_RE

    for row in rows[header_idx + 1 :]:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue

        # ARIAORMS stacks one day per block: a date row, then the same header, then values.
        first_cell = next((cell for cell in row[:3] if cell is not None and str(cell).strip()), None)
        if isinstance(first_cell, datetime):
            reading_date = first_cell.date()
            continue
        if isinstance(first_cell, date) and not isinstance(first_cell, datetime):
            reading_date = first_cell
            continue
        first_text = str(first_cell).strip() if first_cell is not None else ""
        if _date_row.match(first_text):
            parsed_day = parse_plant_date(first_text)
            if parsed_day:
                reading_date = parsed_day
            continue
        header_vals = [str(c).strip().lower() if c is not None else "" for c in row]
        if "parameter" in header_vals and any(re.match(r"^\d{1,2}:\d{2}$", v) for v in header_vals):
            continue

        def _cell(idx: int | None) -> str:
            if idx is None or idx >= len(row) or row[idx] is None:
                return ""
            return str(row[idx]).strip()

        equipment = _cell(equip_col if equip_col is not None else 0)
        tag = _cell(tag_col)
        parameter = _cell(param_col if param_col is not None else (1 if equip_col is None else None))
        label = parameter or tag
        if not label or label.lower() in {"parameter", "tag"}:
            continue

        if not electrolyzer:
            electrolyzer = _detect_electrolyzer(equipment, tag, parameter, label)

        # Extract operator from SiteMan-style cells
        for j in time_cols:
            raw = row[j] if j < len(row) else None
            if isinstance(raw, str) and "\n" in raw:
                op = raw.split("\n", 1)[1].strip()
                if op:
                    operators.add(op)

        cell_m = _CELL_RE.match(label) or _CELL_RE.match(tag) or _CELL_RE.match(parameter)
        if cell_m:
            pos_raw = cell_m.group("pos")
            el = cell_m.group("el").upper()
            electrolyzer = electrolyzer or el
            pos = pos_raw.lstrip("0") or "0"
            for j, tlabel in time_cols.items():
                raw = row[j] if j < len(row) else None
                voltage = parse_float_plant(raw)
                if voltage is None:
                    continue
                # skip absurd totals mistaken as cells
                if voltage > 20:
                    continue
                entry = {
                    "electrolyzer": el,
                    "position": pos,
                    "element_nr": f"{el}-{pos_raw}",
                    "time": tlabel,
                    "voltage": voltage,
                    "date": reading_date.isoformat() if reading_date else None,
                }
                readings.append(entry)
                # Normal Range column often sits just before first time slot
                if time_cols:
                    first_time = min(time_cols.keys())
                    nr_idx = first_time - 1
                    if nr_idx >= 0 and nr_idx < len(row):
                        limit = parse_float_plant(row[nr_idx])
                        # ARIAORMS shows "-3.5" as max delta style; flag high cell voltage
                        if limit is not None and limit < 0 and voltage >= abs(limit):
                            out_of_range.append({**entry, "normal_range": limit})
            continue

        # Aggregate / plant rows (total voltage, temps, load)
        key = label.lower()
        el_from_label = _detect_electrolyzer(label, tag, parameter, equipment)
        if el_from_label:
            electrolyzer = electrolyzer or el_from_label
        day_iso = reading_date.isoformat() if reading_date else ""
        for j, tlabel in time_cols.items():
            raw = row[j] if j < len(row) else None
            voltage = parse_float_plant(raw)
            if voltage is None:
                continue
            slot = total_slots.setdefault(day_iso + "|" + tlabel, {"date": day_iso or None, "time": tlabel})
            if "electrolyzer voltage" in key or "totale voltage" in key or "total voltage" in key:
                if "rack a" in key:
                    slot["rack_a"] = voltage
                elif "rack b" in key:
                    slot["rack_b"] = voltage
                else:
                    slot["total"] = voltage
                totals.setdefault(tlabel, {})
                totals[tlabel].update({k: slot[k] for k in ("rack_a", "rack_b", "total") if k in slot})
            elif "average rack" in key:
                if re.search(r"rack\s*a\b", key):
                    slot["rack_a_avg"] = voltage
                elif re.search(r"rack\s*b\b", key):
                    slot["rack_b_avg"] = voltage
            elif "anolyte temperature" in key:
                slot["anolyte_temp"] = voltage
            elif "catholyte temperature" in key:
                slot["catholyte_temp"] = voltage
            elif "catholyte concentration" in key:
                slot["catholyte_conc"] = voltage
            elif key.startswith("load ") or key.startswith("load["):
                slot["load"] = voltage

    return {
        "sheet": preferred,
        "electrolyzer": electrolyzer,
        "reading_date": reading_date.isoformat() if reading_date else None,
        "operators": sorted(operators),
        "readings": readings,
        "totals": totals,
        "total_slots": list(total_slots.values()),
        "out_of_range": out_of_range,
        "imported_cells": len(readings),
        "times": sorted(set(time_cols.values())),
        "dates": sorted({r["date"] for r in readings if r.get("date")}),
    }


def parse_rectifier_excel(content: bytes) -> dict[str, Any]:
    """Parse Rectifier Room Train log sheets.

    Keeps only DC bus voltage and current (kA). Status and temperature rows
    stay out of the cell-voltage tables.
    """
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    preferred = None
    for name in wb.sheetnames:
        lower = name.lower()
        if "siteman" in lower or "whit" in lower:
            continue
        preferred = name
        break
    if preferred is None:
        preferred = wb.sheetnames[0]
    rows = _sheet_rows(wb, preferred)
    wb.close()

    reading_date: date | None = None
    time_cols: dict[int, str] = {}
    param_col = 1
    header_found = False
    slots: dict[tuple[str, str, str], dict[str, Any]] = {}

    for row in rows:
        if not row or all(c is None or str(c).strip() == "" for c in row):
            continue
        first_text = ""
        for cell in row[:3]:
            if cell is not None and str(cell).strip():
                first_text = str(cell).strip()
                break
        if _DATE_ROW_RE.match(first_text):
            parsed_day = parse_plant_date(first_text)
            if parsed_day:
                reading_date = parsed_day
            continue
        vals = [str(c).strip() if c is not None else "" for c in row]
        vals_lower = [v.lower() for v in vals]
        if "parameter" in vals_lower and any(re.match(r"^\d{1,2}:\d{2}$", v) for v in vals):
            found = _col_index(vals_lower, "parameter")
            if found is not None:
                param_col = found
            time_cols = {j: v for j, v in enumerate(vals) if re.match(r"^\d{1,2}:\d{2}$", v)}
            header_found = bool(time_cols)
            continue
        if not header_found or not time_cols or reading_date is None:
            continue
        if param_col >= len(row) or row[param_col] is None:
            continue
        parameter = str(row[param_col]).strip()
        match = _RECT_ITEM_RE.match(parameter)
        if not match:
            continue
        what = match.group("what").lower()
        if "voltage" in what and "vdc" in what:
            field = "voltage_vdc"
        elif "current" in what and "ka" in what:
            field = "current_ka"
        else:
            continue
        electrolyzer = match.group("el").upper()
        day = reading_date.isoformat()
        for col, tlabel in time_cols.items():
            raw = row[col] if col < len(row) else None
            value = parse_float_plant(raw)
            if value is None:
                continue
            slot = slots.setdefault(
                (electrolyzer, day, tlabel),
                {
                    "electrolyzer": electrolyzer,
                    "date": day,
                    "time": tlabel,
                    "voltage_vdc": None,
                    "current_ka": None,
                },
            )
            slot[field] = value

    if not header_found:
        return {"error": "Header row with Parameter/Tag / time columns not found", "rows": []}
    return {"sheet": preferred, "rows": list(slots.values()), "error": None}


# ---------------------------------------------------------------------------
# 2) Hierarchical lab analysis Excel
# ---------------------------------------------------------------------------

_ANALYTE_KEYS = {"NaOH", "NaCl", "NaClO3", "Ni", "Fe", "NaOCl", "HCl", "pH", "Na2SO4", "Na2CO3"}


def parse_lab_analysis_excel(content: bytes) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    rows = _sheet_rows(wb)
    wb.close()

    samples: list[dict[str, Any]] = []
    ctx: dict[str, Any] = {
        "unit_category": None,
        "unit_tag": None,
        "sc_no": None,
        "fluid": None,
        "sampling": None,
        "remark": None,
    }
    current_params: dict[str, Any] = {}
    current_meta: dict[str, Any] = {}

    def flush():
        nonlocal current_params, current_meta
        if not ctx["sampling"] or not current_params:
            current_params = {}
            current_meta = {}
            return
        fluid = (ctx["fluid"] or "").lower()
        if "caustic" in fluid:
            analysis_type = "caustic_feed"
        elif "anolyt" in fluid:
            analysis_type = "anolyte"
        elif "katholyt" in fluid or "catholyte" in fluid:
            analysis_type = "catholyte"
        elif "brine" in fluid or "sole" in fluid:
            analysis_type = "pure_brine"
        elif "chlor" in fluid:
            analysis_type = "chlorine_gas"
        elif "hcl" in fluid:
            analysis_type = "hcl"
        else:
            analysis_type = "caustic_feed"
        dt = parse_plant_datetime(ctx["sampling"])
        params = dict(current_params)
        params["_meta"] = {
            "unit_category": ctx["unit_category"],
            "unit_tag": ctx["unit_tag"],
            "sc_no": ctx["sc_no"],
            "fluid": ctx["fluid"],
            "remark_supervisor": ctx["remark"],
            "specs": current_meta.get("specs", {}),
            "methods": current_meta.get("methods", {}),
        }
        samples.append(
            {
                "analysis_type": analysis_type,
                "scope": "sub_plant",
                "sub_plant": ctx["unit_tag"] or "Train 1",
                "date": dt.isoformat() if dt else None,
                "time": dt.strftime("%H:%M") if dt else None,
                "parameters": params,
            }
        )
        current_params = {}
        current_meta = {"specs": {}, "methods": {}}

    for row in rows:
        cells = [c for c in row if c is not None and str(c).strip() != ""]
        if not cells:
            continue
        # Flatten row text for context lines
        joined = " | ".join(str(c) for c in row if c is not None)
        joined_clean = re.sub(r"<[^>]+>", "", joined)

        if "Unit Category:" in joined_clean:
            flush()
            ctx["unit_category"] = joined_clean.split("Unit Category:", 1)[1].strip().split("|")[0].strip()
            continue
        if "Unit Tag:" in joined_clean:
            flush()
            ctx["unit_tag"] = joined_clean.split("Unit Tag:", 1)[1].strip().split("|")[0].strip()
            continue
        if "SCNo:" in joined_clean:
            flush()
            sc = joined_clean.split("SCNo:", 1)[1]
            sc = re.sub(r"<[^>]+>", "", sc).strip()
            ctx["sc_no"] = sc.split("|")[0].strip()
            continue
        if "Sampling Time:" in joined_clean:
            flush()
            rest = joined_clean.split("Sampling Time:", 1)[1]
            rest = re.sub(r"<[^>]+>", "", rest)
            # 1405/06/03 08:00 Fluid : Caustic
            m = re.match(r"\s*([0-9./\-: ]+?)\s*(?:Fluid\s*:?\s*(.+))?$", rest.strip(), re.I)
            if m:
                ctx["sampling"] = m.group(1).strip()
                if m.group(2):
                    ctx["fluid"] = m.group(2).strip()
            else:
                ctx["sampling"] = rest.strip()
            current_params = {}
            current_meta = {"specs": {}, "methods": {}}
            continue
        if "Remark Supervisor:" in joined_clean:
            ctx["remark"] = joined_clean.split("Remark Supervisor:", 1)[1].strip()
            continue
        if "Analysis" in joined_clean and "Unit Measurement" in joined_clean:
            continue

        # Analyte row: name in col ~5-6, value, unit, spec, method
        # From sample: blank cols then NaOH, O.S, wt%, 31-33, method
        name = None
        value = None
        unit = None
        spec = None
        method = None
        nonempty = [(i, c) for i, c in enumerate(row) if c is not None and str(c).strip() != ""]
        if len(nonempty) >= 2:
            # find analyte name among known keys
            for i, c in nonempty:
                token = str(c).strip()
                token_key = token.replace(" ", "")
                if token_key in _ANALYTE_KEYS or token in _ANALYTE_KEYS:
                    name = token_key if token_key in _ANALYTE_KEYS else token
                    # next cells
                    rest = [x for j, x in nonempty if j > i]
                    if rest:
                        value = rest[0]
                    if len(rest) > 1:
                        unit = str(rest[1]).strip()
                    if len(rest) > 2:
                        spec = str(rest[2]).strip()
                    if len(rest) > 3:
                        method = str(rest[3]).strip()
                    break
        if name:
            raw_val = str(value).strip() if value is not None else None
            num = parse_float_plant(raw_val) if raw_val and raw_val.upper() not in {"O.S", "OS", "N/A", "ND"} else None
            current_params[name] = num if num is not None else (raw_val or "O.S")
            current_meta.setdefault("specs", {})[name] = spec
            current_meta.setdefault("methods", {})[name] = method
            if unit:
                current_meta.setdefault("units", {})[name] = unit

    flush()
    return {"samples": samples, "imported_samples": len(samples)}


# ---------------------------------------------------------------------------
# 3) Assembly / install / dismantle Excel
# ---------------------------------------------------------------------------

def _ymd_from_row(row: list[Any], day_i: int, month_i: int, year_i: int) -> date | None:
    if year_i >= len(row):
        return None
    y = row[year_i] if year_i < len(row) else None
    m = row[month_i] if month_i < len(row) else None
    d = row[day_i] if day_i < len(row) else None
    if y in (None, "") or m in (None, "") or d in (None, ""):
        return None
    try:
        yi, mi, di = int(y), int(m), int(d)
    except (TypeError, ValueError):
        return parse_plant_date(f"{y}/{m}/{d}")
    if yi <= 0 or mi <= 0 or di <= 0:
        return None
    if yi < 100:
        yi = 1400 + yi if yi <= 80 else 1300 + yi
    try:
        if yi < 1600:
            if not 1300 <= yi <= 1500:
                return None
            return jalali_to_gregorian(yi, mi, di)
        if not 1990 <= yi <= 2045:
            return None
        return date(yi, mi, di)
    except Exception:
        return None


ASSEMBLY_FROM_YEAR = 1395  # Jalali; plant asked to load C90 dates from 95 onward
DENORA_SHELL = "دنوار(ایتالیا)"
BLUESTAR_SHELL = "بلو استار (چین)"
BLUESTAR_ELECTROLYZERS = {"M2", "L2"}
_COATING_SHARED_RE = re.compile(r"^آند\s*و\s*کاتد\s*(.*)$")
_COATING_SPLIT_RE = re.compile(
    r"^(آند|کاتد)\s+(.+?)\s*(?:-|–|ـ|و)\s*(آند|کاتد)\s*(.+)$"
)


def _fold_coating(text: str) -> str:
    text = text.replace("ي", "ی").replace("ك", "ک").replace("کاند", "کاتد")
    text = text.replace("وکاتد", "و کاتد").replace("وآند", "و آند")
    return re.sub(r"\s+", " ", text).strip()


def split_coating_company(value: Any) -> tuple[str | None, str | None]:
    """Map «شرکت پوشش دهنده» onto Anode Coating / Cathode Coating.

    A shared phrase «آند و کاتد تعمیراتی دنورا» becomes «آند تعمیراتی دنورا»
    on the anode and «کاتد تعمیراتی دنورا» on the cathode. Mixed phrases
    (آند دنورا - کاتد آنکان) stay on the matching electrode.
    """
    raw = _clean_code(value)
    if not raw:
        return None, None
    text = _fold_coating(raw)

    def clip(part: str | None) -> str | None:
        part = (part or "").strip(" -/")
        return part[:100] or None

    shared = _COATING_SHARED_RE.match(text)
    if shared and not re.match(r"^(آند|کاتد)\b", shared.group(1).strip()):
        rest = shared.group(1).strip()
        return clip(f"آند {rest}".strip()), clip(f"کاتد {rest}".strip())
    matched = _COATING_SPLIT_RE.match(text)
    if matched and matched.group(1) != matched.group(3):
        parts = {
            matched.group(1): clip(matched.group(2)),
            matched.group(3): clip(matched.group(4)),
        }
        return parts.get("آند"), parts.get("کاتد")
    both = clip(text)
    return both, both


def shell_for_electrolyzer(electrolyzer: str | None) -> str | None:
    if not electrolyzer:
        return None
    code = str(electrolyzer).strip().upper()
    if not code or code in {"0", "NO", "REJ", "―", "-", "—"}:
        return None
    if code in BLUESTAR_ELECTROLYZERS:
        return BLUESTAR_SHELL
    return DENORA_SHELL


def _jalali_year_cell(row: list[Any], ymd: tuple[int, int, int] | None) -> int | None:
    if not ymd:
        return None
    year_i = ymd[2]
    if year_i >= len(row):
        return None
    try:
        year = int(row[year_i])
    except (TypeError, ValueError):
        return None
    if year < 100:
        year = 1400 + year if year <= 80 else 1300 + year
    return year


def _pad_cell_position(value: Any) -> str | None:
    text = _clean_code(value)
    if text and re.fullmatch(r"\d{1,3}", text):
        number = int(text)
        if 1 <= number <= 168:
            return f"{number:03d}"
    return text


def _compact_header(value: Any) -> str:
    """Header match key with spaces removed. Assembly column names rely on this."""
    text = "" if value is None else str(value)
    text = text.replace("\u200c", "").replace("ي", "ی").replace("ك", "ک").replace("آ", "ا").replace("أ", "ا")
    return re.sub(r"\s+", "", text).lower()


def _clean_code(value: Any) -> str | None:
    if value is None or value == "":
        return None
    if isinstance(value, float) and value.is_integer():
        text = str(int(value))
    elif isinstance(value, int) and not isinstance(value, bool):
        text = str(value)
    else:
        text = re.sub(r"\s+", " ", str(value)).strip()
    if not text or text in {"*", "―", "-", "—", "XXXXX", "xxxxx"}:
        return None
    return text


def _assembly_columns(header: list[Any], sub: list[Any] | None) -> dict[str, Any] | None:
    keys = [_compact_header(cell) for cell in header]
    blob = " ".join(keys)
    if "anode" not in blob and "cathode" not in blob:
        return None
    if not any(token in blob for token in ("مونتاژ", "شمارهالمنت", "دیمونتاژ", "نصب", "تاریخ")):
        return None

    def find(pred) -> int | None:
        for index, key in enumerate(keys):
            if key and pred(key):
                return index
        return None

    def ymd_at(start: int | None) -> tuple[int, int, int] | None:
        if start is None:
            return None
        sub_key = _compact_header(sub[start]) if sub and start < len(sub) else ""
        if sub_key in {"روز", "day"}:
            return start, start + 1, start + 2
        return None

    status_cols = [index for index, key in enumerate(keys) if key == "وضعیت"]
    # C90 workshop date is stored as «تاریخ مونتاژ در C90» (user: تاریخ نصب در C90).
    c90_at = find(lambda key: "c90" in key and "تاریخ" in key) or find(
        lambda key: "دیمونتاژ" not in key and "مونتاژ" in key and "تاریخ" in key
    )
    install_at = find(lambda key: "نصب" in key and "تاریخ" in key and "c90" not in key)
    dismantle_at = find(lambda key: "دیمونتاژ" in key or ("دی" in key and "مونتاژ" in key and "تاریخ" in key))
    columns: dict[str, Any] = {
        "anode": find(lambda key: key == "anode" or key.startswith("anode")),
        "cathode": find(lambda key: key == "cathode" or key.startswith("cathode")),
        "membrane_type": find(lambda key: "نوعممبران" in key or "membranetype" in key),
        "membrane_use": find(lambda key: "استفاده" in key or "اسفاده" in key or "ممبراننو" in key),
        "membrane_nr": find(lambda key: "کدممبران" in key or "شمارهوکد" in key),
        "element_nr": find(lambda key: "شمارهالمنت" in key or key in {"elementnr", "element_nr"}),
        "electrolyzer": find(
            lambda key: key in {"el", "electrolyzer", "elektrolyseur"}
            or ("الکترولایزر" in key and "موقعیت" not in key and "تاریخ" not in key)
        ),
        "position": find(lambda key: "موقعیت" in key or key in {"position", "p"}),
        "coating": find(lambda key: "پوشش" in key),
        "status": find(lambda key: "وضعیتفعلی" in key or key in {"وضعیتفعلیالمان", "status"}),
        "notes": find(lambda key: key in {"توضیحات", "ملاحظات"} or "دلیل" in key),
        "assembly_ymd": ymd_at(c90_at),
        "install_ymd": ymd_at(install_at),
        "dismantle": dismantle_at,
        "source": status_cols[0] if status_cols else None,
        "site": status_cols[1] if len(status_cols) > 1 else None,
    }
    if columns["anode"] is None and columns["element_nr"] is None:
        return None
    return columns


def _cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _parse_assembly_sheet(rows: list[list[Any]], *, from_year: int | None = ASSEMBLY_FROM_YEAR) -> list[dict[str, Any]] | None:
    header_at = None
    columns = None
    for index, row in enumerate(rows[:12]):
        sub = rows[index + 1] if index + 1 < len(rows) else None
        found = _assembly_columns(row, sub)
        if found:
            header_at = index
            columns = found
            break
    if columns is None or header_at is None:
        return None

    data_at = header_at + (2 if columns["assembly_ymd"] or columns["install_ymd"] else 1)
    elements: list[dict[str, Any]] = []
    for row in rows[data_at:]:
        if not row:
            continue
        anode = _clean_code(_cell(row, columns["anode"]))
        cathode = _clean_code(_cell(row, columns["cathode"]))
        element_nr = _clean_code(_cell(row, columns["element_nr"]))
        if not anode and not cathode and not element_nr:
            continue
        if from_year is not None:
            c90_year = _jalali_year_cell(row, columns["assembly_ymd"])
            if c90_year is None or c90_year < from_year:
                continue
        membrane_type = _clean_code(_cell(row, columns["membrane_type"]))
        membrane_use = _clean_code(_cell(row, columns["membrane_use"]))
        membrane_nr = _clean_code(_cell(row, columns["membrane_nr"]))
        electrolyzer = _clean_code(_cell(row, columns["electrolyzer"]))
        if electrolyzer:
            electrolyzer = electrolyzer.upper()
        position = _pad_cell_position(_cell(row, columns["position"]))
        anode_coating, cathode_coating = split_coating_company(_cell(row, columns.get("coating")))
        shell = shell_for_electrolyzer(electrolyzer)
        assembly = _ymd_from_row(row, *columns["assembly_ymd"]) if columns["assembly_ymd"] else None
        installed = _ymd_from_row(row, *columns["install_ymd"]) if columns["install_ymd"] else None
        dismantle = parse_plant_date(_cell(row, columns["dismantle"]))
        status = _clean_code(_cell(row, columns.get("status")))
        notes = _clean_code(_cell(row, columns.get("notes")))
        if notes and len(notes) > 200:
            notes = notes[:200]

        elements.append(
            {
                "element_nr": element_nr,
                "anode_nr": anode,
                "cathode_nr": cathode,
                "membrane_type": membrane_type,
                "membrane_nr": membrane_nr,
                "membrane_remark": membrane_use,
                "membrane_info": None,
                "electrolyzer": electrolyzer,
                "position": position,
                "assembly_date": assembly.isoformat() if assembly else None,
                "commissioning_date": installed.isoformat() if installed else None,
                "decommissioning_date": dismantle.isoformat() if dismantle else None,
                "disassembly_date": dismantle.isoformat() if dismantle else None,
                "decommission_reason": notes,
                "remarks": status,
                "anode_coating": anode_coating,
                "cathode_coating": cathode_coating,
                "anode_shell": shell,
                "cathode_shell": shell,
                "coating_split": bool(
                    anode_coating and cathode_coating and anode_coating != cathode_coating
                ),
            }
        )
    return elements


def parse_assembly_excel(content: bytes, *, from_year: int | None = ASSEMBLY_FROM_YEAR) -> dict[str, Any]:
    """Read the plant montage / demontage workbook into assembly-form fields.

    C90 day/month/year (تاریخ مونتاژ در C90) is Assembly Date. The default
    keeps Jalali years from 1395 onward. Pass from_year=None to keep every row.
    Site install is commissioning. Dismantle date fills both decommissioning
    and disassembly.
    """
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    elements: list[dict[str, Any]] = []
    recognized = False
    try:
        for name in wb.sheetnames:
            rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
            parsed = _parse_assembly_sheet(rows, from_year=from_year)
            if parsed is None:
                continue
            recognized = True
            elements.extend(parsed)
    finally:
        wb.close()
    if not recognized:
        return {"elements": [], "recognized": False, "imported": 0}
    if not elements:
        return {
            "elements": [],
            "recognized": True,
            "imported": 0,
            "error": "No assembly rows found. هیچ ردیف مونتاژ یا دمونتاژ در فایل پیدا نشد.",
        }
    return {"elements": elements, "recognized": True, "imported": len(elements), "sheet": "assembly"}


# ---------------------------------------------------------------------------
# 4) TAFKIK segregation Excel
# ---------------------------------------------------------------------------

def _header_key(value: Any) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip().lower())
    return text.replace("ي", "ی").replace("ك", "ک")


def _tafkik_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float):
        if value != value:
            return None
        if value == int(value) and abs(value) < 1e12:
            return str(int(value))
        return format(value, "g")
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    if not text or text in {"*", "―", "-", "—"}:
        return None
    return text


def _tafkik_date_and_note(value: Any) -> tuple[date | None, str | None]:
    parsed = parse_plant_date(value)
    if parsed is not None:
        return parsed, None
    note = _tafkik_text(value)
    if note and re.fullmatch(r"[\d./\-]+", note):
        return None, None
    return None, note


def _tafkik_kind(serial: str) -> str:
    if "کاتد" in serial or "cathode" in serial.lower():
        return "cathode"
    if "آند" in serial or "اند" in serial or "anode" in serial.lower():
        return "anode"
    compact = serial.upper().replace(" ", "")
    if compact.startswith(("UA", "DA", "PA", "LA", "HA")):
        return "anode"
    if compact.startswith(("UC", "DC", "PC", "LC", "HC")):
        return "cathode"
    if compact.startswith("A"):
        return "anode"
    if compact.startswith("C"):
        return "cathode"
    return "unknown"


def _tafkik_columns(header: list[Any], sub: list[Any] | None) -> dict[str, int] | None:
    keys = [_header_key(cell) for cell in header]

    def find(*needles: str, exclude: tuple[str, ...] = ()) -> int | None:
        for index, key in enumerate(keys):
            if exclude and any(ex in key for ex in exclude):
                continue
            if any(needle in key for needle in needles):
                return index
        return None

    serial = find("شماره سریال", "serial", exclude=("زوج", "فرم", "بازرسی"))
    if serial is None:
        serial = find("شماره سریال", "serial")
    if serial is None:
        return None
    cols = {
        "serial": serial,
        "company": find("شرکت", "company"),
        "service_life": find("کارکرد", "service"),
        "install": find("تاریخ نصب", "install"),
        "decommission": find("discomission", "decomission", "decommission", "تاریخ خارج"),
        "dismantle": find("تاریخ دمونتاژ", "dismantle", "demontage", "disassemble"),
        "xrf": find("xrf", exclude=("زوج",)),
        "pair_serial": find("الکترود زوج", "زوج"),
        "pair_xrf": find("xrf الکترود زوج", "xrf زوج"),
        "decomm_v": find("ولتاژدر زمان", "ولتاژ در زمان"),
        "decomm_ka": find("ka در زمان", "kaدر زمان"),
        "decomm_temp": find("دما در زمان"),
        "voltage": find("تفسیر ولتاژ"),
        "warranty": find("گارانتی", "warranty"),
        "coating": find("کیفیت پوشش", "coating"),
        "decision": find("آخرین تصمیم", "decision"),
        "problems": find("مشکل"),
        "remarks": find("تفکیک در حضور"),
        "pallet": find("پالت", "pallet"),
        "inspection_form": find("سریال فرم بازرسی", "فرم بازرسی"),
    }
    # Fallback: bare "ولتاژ" only when تفسیر ولتاژ is missing (legacy Sheet2).
    if cols["voltage"] is None:
        cols["voltage"] = find("ولتاژ", "voltage", exclude=("زمان", "تفسیر"))
    insp = find("تاریخ بازرسی", "inspection")
    if insp is not None:
        cols["insp_day"] = insp
        cols["insp_month"] = insp + 1
        cols["insp_year"] = insp + 2
        if sub:
            for index, key in enumerate(_header_key(cell) for cell in sub):
                if key == "روز":
                    cols["insp_day"] = index
                elif key == "ماه":
                    cols["insp_month"] = index
                elif key == "سال":
                    cols["insp_year"] = index
    return cols


def _tafkik_cell(row: list[Any], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def _tafkik_inspection(row: list[Any], cols: dict[str, int]) -> date | None:
    year = _tafkik_cell(row, cols.get("insp_year"))
    if year is None or str(year).strip() == "":
        return None
    try:
        return jalali_to_gregorian(
            int(year),
            int(_tafkik_cell(row, cols.get("insp_month")) or 1),
            int(_tafkik_cell(row, cols.get("insp_day")) or 1),
        )
    except Exception:
        return None


def _tafkik_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and value == value:
        return float(value)
    text = str(value).strip().replace(",", ".")
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _parse_tafkik_sheet(rows: list[list[Any]]) -> list[dict[str, Any]]:
    header_at = None
    columns = None
    for index, row in enumerate(rows[:8]):
        found = _tafkik_columns(row, rows[index + 1] if index + 1 < len(rows) else None)
        if found:
            header_at = index
            columns = found
            break
    if columns is None or header_at is None:
        return []

    start = header_at + 1
    if start < len(rows):
        sub_keys = {_header_key(cell) for cell in rows[start]}
        if sub_keys & {"روز", "ماه", "سال"}:
            start += 1

    records: list[dict[str, Any]] = []
    for row in rows[start:]:
        serial = _tafkik_text(_tafkik_cell(row, columns["serial"]))
        if not serial:
            continue
        install, install_note = _tafkik_date_and_note(_tafkik_cell(row, columns.get("install")))
        decommission, decommission_note = _tafkik_date_and_note(_tafkik_cell(row, columns.get("decommission")))
        dismantle, dismantle_note = _tafkik_date_and_note(_tafkik_cell(row, columns.get("dismantle")))
        inspector_remarks = _tafkik_text(_tafkik_cell(row, columns.get("remarks")))
        notes = " | ".join(
            part for part in (install_note, decommission_note, dismantle_note) if part
        ) or None
        remarks = inspector_remarks or notes
        if inspector_remarks and notes:
            remarks = f"{inspector_remarks} | {notes}"
        inspection = _tafkik_inspection(row, columns)
        records.append(
            {
                "serial_nr": serial,
                "electrode_kind": _tafkik_kind(serial),
                "company": _tafkik_text(_tafkik_cell(row, columns.get("company"))),
                "service_life": _tafkik_text(_tafkik_cell(row, columns.get("service_life"))),
                "install_date": install.isoformat() if install else None,
                "decommission_date": decommission.isoformat() if decommission else None,
                "disassemble_date": dismantle.isoformat() if dismantle else None,
                "inspection_date": inspection.isoformat() if inspection else None,
                "xrf": _tafkik_text(_tafkik_cell(row, columns.get("xrf"))),
                "pair_serial_nr": _tafkik_text(_tafkik_cell(row, columns.get("pair_serial"))),
                "pair_xrf": _tafkik_text(_tafkik_cell(row, columns.get("pair_xrf"))),
                "decommission_voltage": _tafkik_float(_tafkik_cell(row, columns.get("decomm_v"))),
                "decommission_ka": _tafkik_float(_tafkik_cell(row, columns.get("decomm_ka"))),
                "decommission_temp": _tafkik_float(_tafkik_cell(row, columns.get("decomm_temp"))),
                "voltage_quality": _tafkik_text(_tafkik_cell(row, columns.get("voltage"))),
                "warranty": _tafkik_text(_tafkik_cell(row, columns.get("warranty"))),
                "coating_quality": _tafkik_text(_tafkik_cell(row, columns.get("coating"))),
                "decision": _tafkik_text(_tafkik_cell(row, columns.get("decision"))),
                "problems": _tafkik_text(_tafkik_cell(row, columns.get("problems"))),
                "pallet": _tafkik_text(_tafkik_cell(row, columns.get("pallet"))),
                "inspection_form_serial": _tafkik_text(_tafkik_cell(row, columns.get("inspection_form"))),
                "remarks": remarks,
            }
        )
    return records


def parse_tafkik_excel(content: bytes, sheet: str | None = None) -> dict[str, Any]:
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    names = [sheet] if sheet and sheet in wb.sheetnames else list(wb.sheetnames)
    records: list[dict[str, Any]] = []
    seen: set[tuple] = set()
    previous_sheet_keys: set[tuple] = set()
    used_sheets: list[str] = []
    try:
        for name in names:
            parsed = _parse_tafkik_sheet(_sheet_rows(wb, name))
            if not parsed:
                continue
            used_sheets.append(name)
            sheet_keys: set[tuple] = set()
            for item in parsed:
                identity = (
                    item["serial_nr"],
                    item.get("inspection_date"),
                    item.get("decision"),
                    item.get("problems"),
                    item.get("remarks"),
                    item.get("pallet"),
                )
                serial_insp = (item["serial_nr"], item.get("inspection_date"))
                if identity in seen or serial_insp in previous_sheet_keys:
                    continue
                seen.add(identity)
                sheet_keys.add(serial_insp)
                records.append(item)
            previous_sheet_keys |= sheet_keys
    finally:
        wb.close()
    return {
        "records": records,
        "imported": len(records),
        "sheet": ", ".join(used_sheets),
        "sheets": used_sheets,
    }
