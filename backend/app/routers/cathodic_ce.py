"""Train 1 / Train 2 cathodic current-efficiency sheets."""
from __future__ import annotations

import io
from datetime import date, datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session
from starlette.responses import StreamingResponse

from .. import cathodic_ce as eng
from .. import models
from ..database import get_db

router = APIRouter(prefix="/cathodic-ce", tags=["cathodic-ce"])

YELLOW = PatternFill("solid", fgColor="FFFF00")
PALE = PatternFill("solid", fgColor="FFFDE7")
RESULT = PatternFill("solid", fgColor="E3F2FD")
HEADER = PatternFill("solid", fgColor="FFF59D")
THIN = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


class SheetIn(BaseModel):
    train: int = Field(ge=1, le=2)
    sheet_date: date
    date_label: str | None = None
    total_load: str | None = None
    initial_reading: float | None = None
    final_reading: float | None = None
    test_hours: float | None = None
    readings: dict = Field(default_factory=dict)
    online: dict = Field(default_factory=dict)
    sign_insp_name: str | None = None
    sign_insp_image: str | None = None
    sign_insp_at: str | None = None
    sign_proc_name: str | None = None
    sign_proc_image: str | None = None
    sign_proc_at: str | None = None


def _clean_readings(raw: dict | None) -> dict:
    out: dict[str, dict[str, float]] = {}
    for code, hours in (raw or {}).items():
        if not isinstance(hours, dict):
            continue
        cleaned = {}
        for hour, value in hours.items():
            number = eng._num(value)
            if hour in eng.HOURS and number is not None:
                cleaned[hour] = number
        if cleaned:
            out[str(code)] = cleaned
    return out


def _clean_online(raw: dict | None) -> dict[str, float]:
    out: dict[str, float] = {}
    for letter, value in (raw or {}).items():
        number = eng._num(value)
        if letter in eng.LETTERS and number is not None:
            out[str(letter)] = number
    return out


def _sheet_payload(train: int, sheet_date: date, row: models.CathodicCeSheet | None) -> dict:
    readings = _clean_readings(row.readings if row else None)
    online = _clean_online(row.online if row else None)
    initial = row.initial_reading if row else None
    final = row.final_reading if row else None
    hours = row.test_hours if row else None
    computed = eng.calculate(train, readings, online, initial, final, hours)
    tags = eng.train_tags(train)
    return {
        "train": train,
        "sheet_date": sheet_date.isoformat(),
        "title": f"CELLROOM(Train {train}) EFFICIENCY",
        "saved": row is not None,
        "hours": eng.HOURS,
        "yellow_hours": eng.YELLOW_HOURS,
        "rows": eng.sheet_rows(train),
        "displays": eng.displays(train),
        "production_tag": tags["production"],
        "inputs": {
            "date_label": (row.date_label if row and row.date_label else ""),
            "total_load": (row.total_load if row and row.total_load else ""),
            "initial_reading": initial,
            "final_reading": final,
            "test_hours": hours,
            "readings": readings,
            "online": online,
            "sign_insp_name": (row.sign_insp_name if row and row.sign_insp_name else ""),
            "sign_insp_image": (row.sign_insp_image if row and row.sign_insp_image else ""),
            "sign_insp_at": (row.sign_insp_at if row and row.sign_insp_at else ""),
            "sign_proc_name": (row.sign_proc_name if row and row.sign_proc_name else ""),
            "sign_proc_image": (row.sign_proc_image if row and row.sign_proc_image else ""),
            "sign_proc_at": (row.sign_proc_at if row and row.sign_proc_at else ""),
        },
        "computed": computed,
    }


def _get(db: Session, train: int, sheet_date: date) -> models.CathodicCeSheet | None:
    return (
        db.query(models.CathodicCeSheet)
        .filter(models.CathodicCeSheet.train == train, models.CathodicCeSheet.sheet_date == sheet_date)
        .one_or_none()
    )


def _check_train(train: int) -> int:
    if train not in (1, 2):
        raise HTTPException(status_code=400, detail="Train must be 1 or 2")
    return train


@router.get("/sheet")
def get_sheet(train: int = Query(...), sheet_date: date = Query(...), db: Session = Depends(get_db)):
    train = _check_train(train)
    return _sheet_payload(train, sheet_date, _get(db, train, sheet_date))


@router.put("/sheet")
def save_sheet(body: SheetIn, db: Session = Depends(get_db)):
    train = _check_train(body.train)
    row = _get(db, train, body.sheet_date)
    if row is None:
        row = models.CathodicCeSheet(train=train, sheet_date=body.sheet_date)
        db.add(row)
    row.date_label = (body.date_label or "").strip() or None
    row.total_load = (body.total_load or "").strip() or None
    row.initial_reading = body.initial_reading
    row.final_reading = body.final_reading
    row.test_hours = body.test_hours
    row.readings = _clean_readings(body.readings)
    row.online = _clean_online(body.online)
    row.sign_insp_name = (body.sign_insp_name or "").strip() or None
    row.sign_insp_image = (body.sign_insp_image or "").strip() or None
    row.sign_insp_at = (body.sign_insp_at or "").strip() or None
    row.sign_proc_name = (body.sign_proc_name or "").strip() or None
    row.sign_proc_image = (body.sign_proc_image or "").strip() or None
    row.sign_proc_at = (body.sign_proc_at or "").strip() or None
    row.updated_at = datetime.utcnow()
    db.commit()
    db.refresh(row)
    return _sheet_payload(train, body.sheet_date, row)


def _history_rows(db: Session, train: int) -> list[dict]:
    stored = (
        db.query(models.CathodicCeSheet)
        .filter(models.CathodicCeSheet.train == train)
        .order_by(models.CathodicCeSheet.sheet_date)
        .all()
    )
    out = []
    for row in stored:
        computed = eng.calculate(train, row.readings, row.online, row.initial_reading, row.final_reading, row.test_hours)
        out.append(
            {
                "sheet_date": row.sheet_date.isoformat() if row.sheet_date else "",
                "date_label": row.date_label or "",
                "total_load": row.total_load or "",
                "test_hours": computed["j31"],
                "production": computed["j37"],
                "current_density": computed["j30"],
                "avg_voltage": computed["j33"],
                "e_theoretical": computed["q32"],
                "e_measured": computed["q33"],
                "e_corrected": computed["q34"],
                "power_efficiency": computed["q35"],
                "ce": computed["q37"],
                "extrapolated": computed["q39"],
            }
        )
    return out


@router.get("/history")
def history(train: int = Query(...), db: Session = Depends(get_db)):
    train = _check_train(train)
    return {"train": train, "rows": _history_rows(db, train)}


def _xlsx(name: str, wb: Workbook) -> StreamingResponse:
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


def _write_value(cell, value):
    number = eng._num(value)
    cell.value = number if number is not None else None
    cell.number_format = "0.####"
    cell.border = THIN
    cell.alignment = Alignment(horizontal="center")


@router.get("/sheet.xlsx")
def export_sheet(train: int = Query(...), sheet_date: date = Query(...), db: Session = Depends(get_db)):
    train = _check_train(train)
    row = _get(db, train, sheet_date)
    payload = _sheet_payload(train, sheet_date, row)
    inputs = payload["inputs"]
    readings = inputs["readings"]
    tags = eng.train_tags(train)

    wb = Workbook()
    ws = wb.active
    ws.title = f"Train {train}"
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.freeze_panes = "K4"
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 78
    for col in range(10, 24):
        ws.column_dimensions[get_column_letter(col)].width = 10

    ws.merge_cells("A1:F1")
    ws["A1"] = inputs["date_label"] or sheet_date.isoformat()
    ws.merge_cells("H1:O1")
    ws["H1"] = payload["title"]
    ws["H1"].font = Font(bold=True, size=14)
    ws.merge_cells("R1:V1")
    ws["R1"] = f"Total Load:{inputs['total_load']}" if inputs["total_load"] else "Total Load:"

    ws["A2"] = "I.R"
    ws.merge_cells("B2:I2")
    ws["B2"] = f"100% w/w NaOH Production at initial reading(Mton),{tags['production']}"
    _write_value(ws["J2"], inputs["initial_reading"])
    ws.merge_cells("K2:M2")
    ws["K2"] = "TIME"
    ws["K2"].fill = HEADER
    ws["K2"].font = Font(bold=True)
    ws["K2"].alignment = Alignment(horizontal="center")
    ws.merge_cells("N2:N3")
    ws["N2"] = "on/off"
    ws["N2"].fill = HEADER
    ws["N2"].alignment = Alignment(horizontal="center", vertical="center")

    ws["A3"] = "F.R"
    ws.merge_cells("B3:I3")
    ws["B3"] = f"100% w/w NaOH Production at Final reading(Mton),{tags['production']}"
    _write_value(ws["J3"], inputs["final_reading"])
    for index, hour in enumerate(eng.HOURS):
        cell = ws.cell(3, 11 + index, hour)
        cell.fill = YELLOW if hour in eng.YELLOW_HOURS else HEADER
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
        cell.border = THIN

    excel_row = {
        **{f"L{letter}": 4 + index for index, letter in enumerate(eng.LETTERS)},
        **{f"V{letter}": 16 + index for index, letter in enumerate(eng.LETTERS)},
        "TEMP": 28,
        "CONC": 29,
    }
    for item in payload["rows"]:
        r = excel_row[item["code"]]
        ws.cell(r, 1, item["label"]).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=9)
        ws.cell(r, 2, item["description"])
        hour_map = readings.get(item["code"]) or {}
        for index, hour in enumerate(eng.HOURS):
            cell = ws.cell(r, 11 + index)
            _write_value(cell, hour_map.get(hour))
            if hour in eng.YELLOW_HOURS:
                cell.fill = YELLOW
        avg = ws.cell(r, 10)
        if item["kind"] == "voltage":
            load_r = r - 12
            avg.value = f'=IF(COUNT(K{load_r}:M{load_r})=0,"",SUM(K{r}:M{r})/COUNT(K{r}:M{r}))'
        else:
            avg.value = f'=IF(COUNT(K{r}:M{r})=0,"",SUM(K{r}:M{r})/COUNT(K{r}:M{r}))'
        avg.fill = RESULT
        avg.number_format = "0.####"
        avg.border = THIN
        if item["kind"] == "load":
            flag = ws.cell(r, 14)
            flag.value = f'=IF(COUNT(K{r}:M{r})=0,"",1)'
            flag.fill = PALE
            flag.border = THIN
            flag.alignment = Alignment(horizontal="center")

    ws["A30"] = "A.i"
    ws.merge_cells("B30:I30")
    ws["B30"] = "Avg. Current Density(KA/m2)"
    ws["J30"] = '=IF(COUNT(N4:N15)=0,"",SUM(J4:J15)/(2.72*COUNT(N4:N15)))'
    ws["A31"] = "T.T"
    ws.merge_cells("B31:I31")
    ws["B31"] = "Total Test Running Time(Hours)"
    _write_value(ws["J31"], inputs["test_hours"])
    ws["J31"].fill = YELLOW

    ws["C33"] = "Avg. Voltage Per Electrolyzer(V)"
    ws["J33"] = '=IF(COUNT(N4:N15)=0,"",SUM(J16:J27)/COUNT(N4:N15))'
    ws["C34"] = "Avg. Voltage Per Element(V)"
    ws["J34"] = "=J33/163"
    ws["C35"] = "Avg. Load of all Rectifier(KA)"
    ws["J35"] = '=IF(COUNT(N4:N15)=0,"",SUM(J4:J15)/COUNT(N4:N15))'
    ws["A37"] = "T.P"
    ws.merge_cells("B37:I37")
    ws["B37"] = "Total 100% w/w product NaOH During T.T (ton/day)"
    ws["J37"] = "=(J3-J2)*24/J31"

    results = [
        (32, "Etheoretical", "=(0.1265*J30+2.4)*1000/(1.4923*0.95)", "kwh/t NaOH100%"),
        (33, "Emeasured", "=J33*COUNT(N4:N15)*J35*24/J37", "kwh/t NaOH100%"),
        (34, "Emeasured,corr.", "=Q33-((90-J28)*7*J30/3)", "kwh/t NaOH100%"),
        (35, "Efficiency of Power", "=(Q32/Q34)*100", "%"),
        (37, "CE(Current Efficiency)", "=100*J37*1000/(J35*24*1.4923*163*COUNT(N4:N15))", "%"),
        (39, "Extrapolated Production", "=J37*5.46*12/(J30*COUNT(N4:N15))", "mtpd NaOH 100%"),
    ]
    for r, label, formula, unit in results:
        ws.cell(r, 15, label).font = Font(bold=True)
        ws.merge_cells(start_row=r, start_column=15, end_row=r, end_column=16)
        cell = ws.cell(r, 17, formula)
        cell.fill = RESULT
        cell.font = Font(bold=True, size=12)
        cell.number_format = "0.####"
        ws.cell(r, 20, unit)
    for addr in ("J30", "J33", "J34", "J35", "J37"):
        ws[addr].fill = RESULT
        ws[addr].number_format = "0.####"
        ws[addr].font = Font(bold=True)

    ws["A41"] = "Inspection and process check"
    ws["A41"].font = Font(bold=True)
    ws["B41"] = f"Inspection: {inputs.get('sign_insp_name') or '—'}  {inputs.get('sign_insp_at') or ''}"
    ws["B42"] = f"Process: {inputs.get('sign_proc_name') or '—'}  {inputs.get('sign_proc_at') or ''}"

    return _xlsx(f"cathodic-ce-train{train}-{sheet_date.isoformat()}.xlsx", wb)


@router.get("/history.xlsx")
def export_history(train: int = Query(...), db: Session = Depends(get_db)):
    train = _check_train(train)
    rows = _history_rows(db, train)
    wb = Workbook()
    ws = wb.active
    ws.title = "History"
    headers = [
        "Date",
        "Label",
        "Total load",
        "Test hours",
        "NaOH t/day",
        "Current density KA/m2",
        "Avg voltage V",
        "E theoretical",
        "E measured",
        "E measured corr",
        "Power efficiency %",
        "CE %",
        "Extrapolated mtpd",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER
        cell.font = Font(bold=True)
    keys = [
        "sheet_date",
        "date_label",
        "total_load",
        "test_hours",
        "production",
        "current_density",
        "avg_voltage",
        "e_theoretical",
        "e_measured",
        "e_corrected",
        "power_efficiency",
        "ce",
        "extrapolated",
    ]
    for item in rows:
        ws.append([item[key] for key in keys])
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 18
    if len(rows) >= 1:
        chart = LineChart()
        chart.title = f"Train {train} cathodic CE"
        chart.y_axis.title = "CE %"
        chart.x_axis.title = "Date"
        chart.height = 8
        chart.width = 18
        data = Reference(ws, min_col=12, min_row=1, max_row=len(rows) + 1)
        cats = Reference(ws, min_col=1, min_row=2, max_row=len(rows) + 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)
        chart.shape = 4
        ws.add_chart(chart, "A" + str(len(rows) + 3))
    return _xlsx(f"cathodic-ce-train{train}-history.xlsx", wb)
