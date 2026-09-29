"""Live anode/cathode warehouse.

Stock is not a separate table. Each serial is classified from the latest
assembly installation plus maintenance, recoating and coating-check rows, so a
save in any of those forms shows up here on the next read.
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from io import BytesIO

from openpyxl import load_workbook
from sqlalchemy import func
from sqlalchemy.orm import Session

from . import models
from .calculations import installation_dol

BUCKETS = ("on_rack", "out_repair", "pending", "ok", "not_ok")


def compact_nr(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value)).strip().upper()


def _as_date(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return None


def _iso(value) -> str | None:
    day = _as_date(value)
    return day.isoformat() if day else None


def _stamp(day: date) -> datetime:
    return datetime.combine(day, datetime.min.time())


def _fold(text: str | None) -> str:
    return (text or "").replace("ي", "ی").replace("ك", "ک").replace("ة", "ه")


def load_compact_index(db: Session, model, field: str) -> dict[str, str]:
    index: dict[str, str] = {}
    column = getattr(model, field)
    for (nr,) in db.query(column).all():
        key = compact_nr(nr)
        if key and key not in index:
            index[key] = nr
    return index


def ensure_indexed(db: Session, model, field: str, raw, index: dict[str, str] | None, extra: dict | None = None) -> int:
    """Create a master row when this serial is not already in the catalog. Does not overwrite details."""
    text = (raw or "").strip()
    if not text or text.upper() in {"*", "-", "—", "―", "XXXXX"}:
        return 0
    key = compact_nr(text)
    if not key:
        return 0
    if index is not None and key in index:
        return 0
    column = getattr(model, field)
    if index is None:
        hit = db.query(model).filter(func.replace(func.upper(column), " ", "") == key).first()
        if hit:
            return 0
    stored = text[:50]
    kwargs = {field: stored}
    if extra:
        for name, value in extra.items():
            if value:
                kwargs[name] = str(value)[:100]
    db.add(model(**kwargs))
    if index is not None:
        index[key] = stored
    return 1


def _sent_out(text: str | None) -> bool:
    folded = _fold(text)
    # "ارسال به" is sent out to a contractor. "ارسالی از" is where the part came from.
    return "ارسال به" in folded or "sent to" in folded.lower()


def _damaged(kind: str, text: str | None) -> bool:
    folded = _fold(text)
    if not folded.strip():
        return False
    low = folded.lower()
    if ("ولتاژ" in folded or "ولتاز" in folded) and "بالا" in folded:
        return True
    if any(token in folded for token in ("نشتی", "برآمدگی", "معیوب")):
        return True
    if "not ok" in low or "notok" in low:
        return True
    hole = "سوراخ" in folded or "hole" in low
    if not hole:
        return False
    membrane = "ممبر" in folded or "membrane" in low
    anode = "آند" in folded or "anode" in low
    cathode = "کاتد" in folded or "cathode" in low
    if membrane and not anode and not cathode:
        return False
    if kind == "anode":
        if anode:
            return True
        if cathode:
            return False
        return True
    if cathode:
        return True
    if anode:
        return False
    return True


def _verdict(text: str | None) -> str | None:
    low = _fold(text).lower()
    if "not ok" in low or "notok" in low:
        return "not_ok"
    if re.search(r"\bok\b", low):
        return "ok"
    return None


def _life_key(el: models.Element):
    return (
        _as_date(el.assembly_date) or _as_date(el.commissioning_date) or date.min,
        _as_date(el.commissioning_date) or date.min,
        el.id or 0,
    )


def _grouped(rows, field: str) -> dict[str, list]:
    grouped: dict[str, list] = defaultdict(list)
    for row in rows:
        key = compact_nr(getattr(row, field))
        if key:
            grouped[key].append(row)
    return grouped


def _open_job(records, after: date | None):
    found = []
    for rec in records:
        dispatch = _as_date(getattr(rec, "dispatch_date", None))
        returned = _as_date(getattr(rec, "return_date", None))
        if dispatch and not returned and (after is None or dispatch >= after):
            found.append(rec)
    if not found:
        return None
    return max(found, key=lambda rec: (_as_date(rec.dispatch_date) or date.min, rec.id or 0))


def _returned_job(records, after: date | None):
    found = []
    for rec in records:
        returned = _as_date(getattr(rec, "return_date", None))
        if returned and (after is None or returned >= after):
            found.append(rec)
    if not found:
        return None
    return max(found, key=lambda rec: (_as_date(rec.return_date) or date.min, rec.id or 0))


def _verdict_check(checks, after: date | None):
    best = None
    best_key = None
    for check in checks:
        day = _as_date(check.check_date)
        if after and (day is None or day < after):
            continue
        verdict = _verdict(check.remarks)
        if not verdict:
            continue
        key = (day or date.min, check.id or 0)
        if best is None or key >= best_key:
            best = check
            best_key = key
    return best


def _segregation_bucket(rows, after: date | None) -> str | None:
    relevant = []
    for row in rows:
        day = _as_date(row.inspection_date) or _as_date(row.dismantle_date)
        if after and day and day < after:
            continue
        relevant.append(row)
    if not relevant:
        return None
    latest = max(relevant, key=lambda row: (_as_date(row.inspection_date) or date.min, row.id or 0))
    text = _fold(f"{latest.decision or ''} {latest.problems or ''} {latest.segregation or ''}")
    low = text.lower()
    if "not ok" in low or "اسقاط" in text or "معیوب" in text:
        return "not_ok"
    if "ارسال" in text:
        return "out_repair"
    if re.search(r"\bok\b", low):
        return "ok"
    return None


_KIND = {
    "anode": {
        "field": "anode_nr",
        "catalog": models.Anode,
        "maintenance": models.AnodeMaintenance,
        "recoating": models.AnodeRecoating,
        "check": models.AnodeCoatingCheck,
    },
    "cathode": {
        "field": "cathode_nr",
        "catalog": models.Cathode,
        "maintenance": models.CathodeMaintenance,
        "recoating": models.CathodeRecoating,
        "check": models.CathodeCoatingCheck,
    },
}


def _classify(kind: str, latest, maintenance, recoating, checks, segs, catalog) -> tuple[str, str, object | None, object | None]:
    after = _as_date(latest.disassembly_date) if latest is not None else None
    mounted = latest is not None and after is None and (
        latest.electrolyzer or latest.commissioning_date or latest.assembly_date
    )
    if mounted:
        return "on_rack", "in_service", None, None

    open_maint = _open_job(maintenance, after)
    open_coat = _open_job(recoating, after)
    if open_maint or open_coat:
        chosen = open_maint or open_coat
        if open_maint and open_coat:
            chosen = max(
                (open_maint, open_coat),
                key=lambda rec: (_as_date(rec.dispatch_date) or date.min, rec.id or 0),
            )
        reason = "maintenance" if chosen is open_maint else "recoating"
        return "out_repair", reason, chosen, None

    ret_maint = _returned_job(maintenance, after)
    ret_coat = _returned_job(recoating, after)
    returned = None
    if ret_maint or ret_coat:
        returned = max(
            [rec for rec in (ret_maint, ret_coat) if rec],
            key=lambda rec: (_as_date(rec.return_date) or date.min, rec.id or 0),
        )
    verdict_after = _as_date(returned.return_date) if returned is not None else after
    verdict = _verdict_check(checks, verdict_after)
    if verdict is not None:
        return _verdict(verdict.remarks) or "ok", "coating", returned, verdict

    remarks = latest.remarks if latest is not None else None
    if latest is not None and _sent_out(remarks) and returned is None:
        return "out_repair", "remarks", None, None
    if returned is not None:
        return "pending", "returned", returned, None
    seg = _segregation_bucket(segs, after)
    if seg:
        return seg, "segregation", None, None
    if latest is not None and _damaged(kind, remarks):
        return "not_ok", "damage", None, None
    if catalog is not None and _as_date(catalog.decommission_date):
        return "not_ok", "decommission", None, None
    return "ok", "ready", returned, None


def build_board(db: Session, kind: str, bucket: str | None = None, q: str | None = None, limit: int = 2500) -> dict:
    spec = _KIND[kind]
    field = spec["field"]
    elements = _grouped(db.query(models.Element).all(), field)
    catalogs = {compact_nr(getattr(row, field)): row for row in db.query(spec["catalog"]).all() if compact_nr(getattr(row, field))}
    maintenance = _grouped(db.query(spec["maintenance"]).all(), field)
    recoating = _grouped(db.query(spec["recoating"]).all(), field)
    checks = _grouped(db.query(spec["check"]).all(), field)
    segs: dict[str, list] = defaultdict(list)
    for row in db.query(models.ElectrodeSegregation).all():
        if row.electrode_kind and row.electrode_kind not in {kind, "unknown"}:
            continue
        key = compact_nr(row.serial_nr)
        if key:
            segs[key].append(row)

    keys = set(elements) | set(catalogs) | set(maintenance) | set(recoating) | set(checks)
    counts = {name: 0 for name in BUCKETS}
    items = []
    for key in keys:
        lives = elements.get(key) or []
        latest = max(lives, key=_life_key) if lives else None
        catalog = catalogs.get(key)
        bucket_name, reason, repair, coating = _classify(
            kind,
            latest,
            maintenance.get(key) or [],
            recoating.get(key) or [],
            checks.get(key) or [],
            segs.get(key) or [],
            catalog,
        )
        counts[bucket_name] += 1
        display = None
        if latest is not None:
            display = getattr(latest, field)
        if not display and catalog is not None:
            display = getattr(catalog, field)
        display = (display or key).strip()
        dols = [
            installation_dol(el.assembly_date, el.commissioning_date, el.disassembly_date, el.decommissioning_date)
            for el in lives
        ]
        dols = [value for value in dols if value is not None]
        last_dol = None
        if latest is not None:
            last_dol = installation_dol(
                latest.assembly_date, latest.commissioning_date, latest.disassembly_date, latest.decommissioning_date
            )
        coating_text = None
        if coating is not None:
            coating_text = coating.remarks
        elif checks.get(key):
            newest = max(checks[key], key=lambda row: (_as_date(row.check_date) or date.min, row.id or 0))
            coating_text = newest.remarks
        items.append(
            {
                "serial": display,
                "key": key,
                "kind": kind,
                "bucket": bucket_name,
                "reason": reason,
                "electrolyzer": latest.electrolyzer if latest else None,
                "position": latest.position if latest else None,
                "element_nr": latest.element_nr if latest else None,
                "assembly_date": _iso(latest.assembly_date) if latest else None,
                "commissioning_date": _iso(latest.commissioning_date) if latest else None,
                "disassembly_date": _iso(latest.disassembly_date) if latest else None,
                "last_dol": last_dol,
                "total_dol": sum(dols) if dols else None,
                "runs": len(lives),
                "remarks": latest.remarks if latest else None,
                "repair": reason if reason in {"maintenance", "recoating", "returned"} else None,
                "repair_dispatch": _iso(getattr(repair, "dispatch_date", None)) if repair is not None else None,
                "repair_return": _iso(getattr(repair, "return_date", None)) if repair is not None else None,
                "coating": coating_text,
                "manufacturer": catalog.manufacturer if catalog is not None else None,
            }
        )

    if bucket == "warehouse":
        items = [item for item in items if item["bucket"] != "on_rack"]
    elif bucket and bucket != "all":
        items = [item for item in items if item["bucket"] == bucket]
    if q:
        needle = q.strip().lower()
        compact_needle = compact_nr(q)
        items = [
            item
            for item in items
            if (compact_needle and compact_needle in item["key"])
            or needle in (item["serial"] or "").lower()
            or needle in (item["remarks"] or "").lower()
            or needle in (item["electrolyzer"] or "").lower()
            or needle in (item["element_nr"] or "").lower()
            or needle in (item["coating"] or "").lower()
        ]
    order = {"out_repair": 0, "pending": 1, "not_ok": 2, "ok": 3, "on_rack": 4}
    items.sort(key=lambda item: (order.get(item["bucket"], 9), item["serial"]))
    truncated = len(items) > limit
    return {
        "kind": kind,
        "counts": {**counts, "total": sum(counts.values())},
        "items": items[:limit],
        "truncated": truncated,
    }


def _rows_for(db: Session, model, field: str, serial: str):
    key = compact_nr(serial)
    column = getattr(model, field)
    return db.query(model).filter(func.replace(func.upper(column), " ", "") == key).all()


def _close_open(rows, day: date) -> int:
    closed = 0
    for row in rows:
        if _as_date(getattr(row, "dispatch_date", None)) and not _as_date(getattr(row, "return_date", None)):
            row.return_date = _stamp(day)
            closed += 1
    return closed


def apply_move(db: Session, kind: str, serial: str, action: str, when: date | None, note: str | None) -> dict:
    spec = _KIND.get(kind)
    if spec is None:
        raise ValueError("kind must be anode or cathode")
    actions = {"dispatch_maintenance", "dispatch_recoating", "return", "ok", "not_ok"}
    if action not in actions:
        raise ValueError("unknown warehouse action")
    stored = (serial or "").strip()[:50]
    if not stored:
        raise ValueError("serial is required")
    day = when or date.today()
    field = spec["field"]
    ensure_indexed(db, spec["catalog"], field, stored, None)
    maintenance = _rows_for(db, spec["maintenance"], field, stored)
    recoating = _rows_for(db, spec["recoating"], field, stored)
    note_text = (note or "").strip() or None

    if action == "dispatch_maintenance":
        if _open_job(maintenance, None):
            return {"ok": True, "already": True, "serial": stored, "action": action}
        db.add(
            spec["maintenance"](
                **{
                    field: stored,
                    "date": _stamp(day),
                    "finding": note_text or "ارسال برای تعمیر",
                    "action": "تعمیر",
                    "dispatch_date": _stamp(day),
                }
            )
        )
    elif action == "dispatch_recoating":
        if _open_job(recoating, None):
            return {"ok": True, "already": True, "serial": stored, "action": action}
        payload = {
            field: stored,
            "dispatch_date": _stamp(day),
            "remarks": note_text or "ارسال برای بازپوشش",
        }
        db.add(spec["recoating"](**payload))
    elif action == "return":
        closed = _close_open(maintenance, day) + _close_open(recoating, day)
        if closed == 0:
            db.add(
                spec["maintenance"](
                    **{
                        field: stored,
                        "date": _stamp(day),
                        "finding": note_text or "بازگشت به انبار",
                        "action": "بازگشت",
                        "dispatch_date": _stamp(day),
                        "return_date": _stamp(day),
                    }
                )
            )
    else:
        _close_open(maintenance, day)
        _close_open(recoating, day)
        lives = _rows_for(db, models.Element, field, stored)
        latest = max(lives, key=_life_key) if lives else None
        dol = None
        if latest is not None:
            dol = installation_dol(
                latest.assembly_date, latest.commissioning_date, latest.disassembly_date, latest.decommissioning_date
            )
        label = "OK" if action == "ok" else "NOT OK"
        remarks = f"{label} — {note_text}" if note_text else label
        payload = {field: stored, "check_date": _stamp(day), "remarks": remarks}
        if kind == "anode":
            payload["dol_days"] = dol
        db.add(spec["check"](**payload))

    db.commit()
    return {"ok": True, "serial": stored, "action": action, "date": day.isoformat()}


def _match_serial(index: dict[str, list], token) -> str | None:
    key = compact_nr(token)
    if not key:
        return None
    if key in index and index[key]:
        return key
    hits = [name for name in index if name.endswith(key) and len(name) > len(key)]
    if len(hits) == 1:
        return hits[0]
    return None


def import_contractor_repairs(db: Session, content: bytes) -> int:
    """The contractor sheet (Anode / Cathode / EL / P) becomes recoating rows."""
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    created = 0
    try:
        anode_index = _grouped(db.query(models.Element).filter(models.Element.anode_nr.isnot(None)).all(), "anode_nr")
        cathode_index = _grouped(db.query(models.Element).filter(models.Element.cathode_nr.isnot(None)).all(), "cathode_nr")
        for name in wb.sheetnames:
            rows = [list(row) for row in wb[name].iter_rows(values_only=True)]
            if not rows:
                continue
            header = [compact_nr(cell) for cell in rows[0][:4]]
            if header[:2] != ["ANODE", "CATHODE"] or "EL" not in header:
                continue
            manufacturer = str(name).strip()[:100] or None
            for row in rows[1:]:
                if not row or row[0] in (None, ""):
                    continue
                anode_token = row[0]
                cathode_token = row[1] if len(row) > 1 else None
                electrolyzer = str(row[2]).strip() if len(row) > 2 and row[2] not in (None, "") else None
                position = str(row[3]).strip() if len(row) > 3 and row[3] not in (None, "") else None
                if isinstance(position, str) and position.endswith(".0"):
                    position = position[:-2]
                created += _add_contractor(
                    db, "anode", anode_token, anode_index, manufacturer, electrolyzer, position
                )
                created += _add_contractor(
                    db, "cathode", cathode_token, cathode_index, manufacturer, electrolyzer, position
                )
    finally:
        wb.close()
    if created:
        db.commit()
    return created


def _add_contractor(db: Session, kind: str, token, index: dict[str, list], manufacturer, electrolyzer, position) -> int:
    if token in (None, ""):
        return 0
    spec = _KIND[kind]
    field = spec["field"]
    matched = _match_serial(index, token)
    lives = index.get(matched) or [] if matched else []
    latest = max(lives, key=_life_key) if lives else None
    if electrolyzer and position and lives:
        exact = [
            el
            for el in lives
            if compact_nr(el.electrolyzer) == compact_nr(electrolyzer) and str(el.position or "").strip() == str(position).strip()
        ]
        if exact:
            latest = max(exact, key=_life_key)
    serial = getattr(latest, field) if latest is not None else str(token).strip()[:50]
    if not serial:
        return 0
    ensure_indexed(db, spec["catalog"], field, serial, None)
    existing = _rows_for(db, spec["recoating"], field, serial)
    if any((row.manufacturer or "").strip() == (manufacturer or "") for row in existing):
        return 0
    on_rack = latest is not None and latest.disassembly_date is None
    dispatch = None
    returned = None
    if latest is not None:
        dispatch = _as_date(latest.disassembly_date) or _as_date(latest.assembly_date)
        if on_rack:
            returned = _as_date(latest.commissioning_date) or dispatch
    remarks = " ".join(part for part in (f"EL {electrolyzer}" if electrolyzer else "", f"P {position}" if position else "") if part)
    db.add(
        spec["recoating"](
            **{
                field: str(serial).strip()[:50],
                "dispatch_date": _stamp(dispatch) if dispatch else None,
                "return_date": _stamp(returned) if returned else None,
                "manufacturer": manufacturer,
                "remarks": remarks or "ارسال برای بازپوشش",
            }
        )
    )
    return 1
