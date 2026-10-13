"""Recoating pallet sheets (one sheet per CO code) and reading back the sheets the coating company filled in."""
from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["No.", "Type", "Element No.", "Items / Problems", "Pre inspection", "Final Report", "Remark"]
CODE_LABEL = "Code"
_THIN = Side(style="thin", color="7F7F7F")
_BOX = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_HEAD_FILL = PatternFill("solid", fgColor="DCE6F1")


def _sheet_title(code: str) -> str:
    return code.replace("/", "-")[:31]


def build_workbook(sheets: list[dict[str, Any]]) -> bytes:
    """sheets: [{code, pallet_no, pallets, date, company, indent_no, items: [{kind, serial, problem}]}]."""
    wb = Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        ws = wb.create_sheet(_sheet_title(spec["code"]))
        ws.sheet_view.rightToLeft = False
        ws["A1"] = "Anode / Cathode Recoating Report"
        ws["A1"].font = Font(bold=True, size=14)
        ws.merge_cells("A1:G1")
        info = [
            (CODE_LABEL, spec["code"]),
            ("Date", spec.get("date") or ""),
            ("Pallet", f"{spec.get('pallet_no', 1)} / {spec.get('pallets', 1)}"),
            ("Coating company", spec.get("company") or ""),
            ("Indent No.", spec.get("indent_no") or ""),
        ]
        for i, (label, value) in enumerate(info, start=2):
            ws.cell(row=i, column=1, value=label).font = Font(bold=True)
            ws.cell(row=i, column=2, value=value)
        head_row = len(info) + 3
        for col, title in enumerate(HEADERS, start=1):
            c = ws.cell(row=head_row, column=col, value=title)
            c.font = Font(bold=True)
            c.fill = _HEAD_FILL
            c.border = _BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        row = head_row
        for n, item in enumerate(spec["items"], start=1):
            row = head_row + n
            values = [n, item["kind"].capitalize(), item["serial"], item.get("problem") or "", "", "", ""]
            for col, value in enumerate(values, start=1):
                c = ws.cell(row=row, column=col, value=value)
                c.border = _BOX
                c.alignment = Alignment(vertical="top", wrap_text=True)
        sign = row + 2
        for col, title in ((5, "Pre inspection"), (6, "Final Report")):
            ws.cell(row=sign, column=col, value=title).font = Font(bold=True)
            for k, label in enumerate(("Name:", "Date:", "Signature:"), start=1):
                c = ws.cell(row=sign + k, column=col, value=label)
                c.border = _BOX
        for col, width in enumerate((6, 10, 18, 34, 26, 26, 24), start=1):
            ws.column_dimensions[get_column_letter(col)].width = width
        ws.freeze_panes = ws.cell(row=head_row + 1, column=1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_report(content: bytes) -> list[dict[str, Any]]:
    """Return [{code | None, rows: [{serial, pre_inspection, final_report, remark}]}] for every sheet with the item table."""
    wb = load_workbook(BytesIO(content), data_only=True, read_only=True)
    out: list[dict[str, Any]] = []
    for ws in wb.worksheets:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        code = None
        head_idx = None
        cols: dict[str, int] = {}
        for i, r in enumerate(rows[:40]):
            cells = [_text(v) for v in r]
            if code is None and cells and cells[0].lower() == CODE_LABEL.lower() and len(cells) > 1:
                code = cells[1] or None
            lowered = [c.lower() for c in cells]
            if "element no." in lowered:
                head_idx = i
                for name in ("element no.", "pre inspection", "final report", "remark"):
                    if name in lowered:
                        cols[name] = lowered.index(name)
                break
        if head_idx is None:
            continue
        parsed: list[dict[str, str]] = []
        for r in rows[head_idx + 1 :]:
            cells = [_text(v) for v in r]

            def at(name: str) -> str:
                idx = cols.get(name)
                return cells[idx] if idx is not None and idx < len(cells) else ""

            serial = at("element no.")
            if not serial:
                if any(cells):
                    continue
                break
            parsed.append(
                {
                    "serial": serial,
                    "pre_inspection": at("pre inspection"),
                    "final_report": at("final report"),
                    "remark": at("remark"),
                }
            )
        out.append({"code": code, "rows": parsed})
    wb.close()
    return out
