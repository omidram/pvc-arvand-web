"""Build Word (.docx) CEO work report from the Persian markdown source (true RTL)."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
MD = DOCS / "گزارش-کار-کامل-پروژه-Arvand-Electrolyzer.md"
OUT = DOCS / "Arvand-Electrolyzer-CEO-Work-Report.docx"
OUT_FA = DOCS / "گزارش-کار-کامل-پروژه-Arvand-Electrolyzer.docx"
LOGO = ROOT / "frontend" / "public" / "logo-arvand.png"

NAVY = RGBColor(0x0B, 0x1F, 0x4A)
GOLD = RGBColor(0x8B, 0x7A, 0x12)
SLATE = RGBColor(0x33, 0x41, 0x55)
MUTED = RGBColor(0x64, 0x74, 0x8B)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Tahoma"


def _ensure_child(parent, tag: str):
    el = parent.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        parent.append(el)
    return el


def set_run_font(run, *, bold=False, size=11, color=SLATE, name=FONT, rtl=True):
    """Apply Persian-friendly fonts + complex-script RTL on the run."""
    r = run._element
    rPr = r.get_or_add_rPr()

    rFonts = _ensure_child(rPr, "w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), name)

    # Critical for Word RTL: mark the run itself as RTL
    if rtl:
        rtl_el = _ensure_child(rPr, "w:rtl")
        rtl_el.set(qn("w:val"), "1")

    sz = _ensure_child(rPr, "w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    szCs = _ensure_child(rPr, "w:szCs")
    szCs.set(qn("w:val"), str(int(size * 2)))

    if bold:
        _ensure_child(rPr, "w:b")
        _ensure_child(rPr, "w:bCs")

    color_el = _ensure_child(rPr, "w:color")
    color_el.set(qn("w:val"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")

    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def set_rtl(paragraph, *, align=WD_ALIGN_PARAGRAPH.RIGHT):
    """Paragraph-level RTL (bidi + right/center alignment)."""
    pPr = paragraph._p.get_or_add_pPr()

    bidi = _ensure_child(pPr, "w:bidi")
    bidi.set(qn("w:val"), "1")

    # Keep text flow RTL even when visually centered (cover)
    jc = _ensure_child(pPr, "w:jc")
    if align == WD_ALIGN_PARAGRAPH.CENTER:
        jc.set(qn("w:val"), "center")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == WD_ALIGN_PARAGRAPH.JUSTIFY:
        jc.set(qn("w:val"), "both")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    else:
        jc.set(qn("w:val"), "right")
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE


def shade_cell(cell, hex_color: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = _ensure_child(tcPr, "w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")


def set_table_rtl(table):
    """Make Word render the table visually RTL (first logical col on the right)."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    bidi = _ensure_child(tblPr, "w:bidiVisual")
    bidi.set(qn("w:val"), "1")


def set_cell_rtl(cell):
    # Cell-level RTL
    tcPr = cell._tc.get_or_add_tcPr()
    tcBidi = _ensure_child(tcPr, "w:textDirection")
    # keep default horizontal; paragraph bidi handles Persian
    for p in cell.paragraphs:
        set_rtl(p)
        for run in p.runs:
            set_run_font(run, size=9, rtl=True)


def configure_document_rtl(doc: Document):
    """Document defaults: Normal style RTL, fa-IR theme language, section bidi."""
    # Normal style
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    pPr = normal.element.get_or_add_pPr()
    bidi = _ensure_child(pPr, "w:bidi")
    bidi.set(qn("w:val"), "1")
    jc = _ensure_child(pPr, "w:jc")
    jc.set(qn("w:val"), "right")
    rPr = normal.element.get_or_add_rPr()
    rFonts = _ensure_child(rPr, "w:rFonts")
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(attr), FONT)
    rtl = _ensure_child(rPr, "w:rtl")
    rtl.set(qn("w:val"), "1")

    # settings.xml — Persian as bidi language
    settings = doc.settings.element
    lang = _ensure_child(settings, "w:themeFontLang")
    lang.set(qn("w:val"), "en-US")
    lang.set(qn("w:eastAsia"), "en-US")
    lang.set(qn("w:bidi"), "fa-IR")

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.6)
    section.bottom_margin = Cm(1.6)

    sectPr = section._sectPr
    sect_bidi = _ensure_child(sectPr, "w:bidi")
    sect_bidi.set(qn("w:val"), "1")


def add_para(
    doc,
    text: str,
    *,
    bold=False,
    size=11,
    color=SLATE,
    space_after=6,
    align=WD_ALIGN_PARAGRAPH.RIGHT,
):
    p = doc.add_paragraph()
    set_rtl(p, align=align)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, bold=bold, size=size, color=color, rtl=True)
    return p


def add_heading_banner(doc, text: str, level: int = 1):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run(text)
    set_run_font(run, bold=True, size=14 if level == 1 else 12, color=NAVY, rtl=True)
    pPr = p._p.get_or_add_pPr()
    shd = _ensure_child(pPr, "w:shd")
    shd.set(qn("w:fill"), "EEF2FF" if level == 1 else "FFFBEB")
    shd.set(qn("w:val"), "clear")
    return p


def add_table(doc, headers: list[str], rows: list[list[str]]):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_rtl(table)

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        set_rtl(p)
        run = p.add_run(h)
        set_run_font(run, bold=True, size=9, color=WHITE, rtl=True)
        shade_cell(cell, "0B1F4A")

    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            p = cell.paragraphs[0]
            set_rtl(p)
            run = p.add_run(val)
            set_run_font(run, bold=False, size=9, color=SLATE, rtl=True)
            if r_idx % 2 == 1:
                shade_cell(cell, "F8FAFC")
    doc.add_paragraph()


def parse_md_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    header = [c.strip() for c in lines[start].strip().strip("|").split("|")]
    i = start + 1
    if i < len(lines) and re.match(r"^\s*\|?\s*[-:| ]+\|?\s*$", lines[i]):
        i += 1
    rows = []
    while i < len(lines) and lines[i].strip().startswith("|"):
        row = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if len(row) < len(header):
            row += [""] * (len(header) - len(row))
        rows.append(row[: len(header)])
        i += 1
    return header, rows, i


def clean_md(text: str) -> str:
    text = text.replace("**", "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    return text.strip()


def build():
    if not MD.exists():
        raise SystemExit(f"Missing markdown source: {MD}")

    lines = MD.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document_rtl(doc)

    if LOGO.exists():
        p = doc.add_paragraph()
        set_rtl(p, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run()
        run.add_picture(str(LOGO), width=Cm(6.5))

    # Cover strip (centered, but Persian runs still RTL)
    cover = doc.add_paragraph()
    set_rtl(cover, align=WD_ALIGN_PARAGRAPH.CENTER)
    pPr = cover._p.get_or_add_pPr()
    shd = _ensure_child(pPr, "w:shd")
    shd.set(qn("w:fill"), "0B1F4A")
    shd.set(qn("w:val"), "clear")
    r1 = cover.add_run("گزارش کار کامل پروژه\n")
    set_run_font(r1, bold=True, size=12, color=GOLD, rtl=True)
    r2 = cover.add_run("Arvand Electrolyzer Management Program\n")
    set_run_font(r2, bold=True, size=16, color=WHITE, rtl=False)
    r3 = cover.add_run("سامانه یکپارچه مدیریت الکترولایزر — پتروشیمی اروند")
    set_run_font(r3, bold=False, size=11, color=RGBColor(0xE2, 0xE8, 0xF0), rtl=True)

    add_para(
        doc,
        "تهیه‌کننده: امید رامشینی (Omid Ramshini) — شرکت وفاق شریف — پشتیبانی فنی: 09398637969",
        bold=True,
        size=10,
        color=NAVY,
        space_after=10,
    )

    i = 0
    skipped_titles = 0
    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        if not line.strip():
            i += 1
            continue

        if re.match(r"^\s*-{3,}\s*$", line):
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m and skipped_titles < 3 and m.group(1) in ("#", "##", "###"):
            # Skip MD cover titles already drawn on the Word cover
            skipped_titles += 1
            i += 1
            continue

        if m:
            level = len(m.group(1))
            text = clean_md(m.group(2))
            if level <= 2:
                add_heading_banner(doc, text, level=1 if level == 1 else 2)
            else:
                add_para(doc, text, bold=True, size=11, color=NAVY, space_after=4)
            i += 1
            continue

        if line.strip().startswith("|") and i + 1 < len(lines):
            headers, rows, nxt = parse_md_table(lines, i)
            # Natural MD order + w:bidiVisual → first column appears on the right
            add_table(
                doc,
                [clean_md(h) for h in headers],
                [[clean_md(c) for c in r] for r in rows],
            )
            i = nxt
            continue

        m = re.match(r"^\s*[-*]\s+\[([ xX])\]\s+(.*)$", line)
        if m:
            done = m.group(1).lower() == "x"
            mark = "✓" if done else "○"
            # With w:rtl, first glyphs sit on the right — mark first is correct
            add_para(doc, f"{mark}  {clean_md(m.group(2))}", size=10, color=SLATE, space_after=2)
            i += 1
            continue

        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            add_para(doc, f"•  {clean_md(m.group(1))}", size=10, color=SLATE, space_after=2)
            i += 1
            continue

        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            add_para(doc, f"{m.group(1)}.  {clean_md(m.group(2))}", size=10, color=SLATE, space_after=2)
            i += 1
            continue

        chunk = [clean_md(line)]
        j = i + 1
        while j < len(lines):
            nxt = lines[j].rstrip()
            if not nxt.strip():
                break
            if (
                nxt.strip().startswith("|")
                or re.match(r"^#{1,4}\s+", nxt)
                or re.match(r"^\s*[-*]\s+", nxt)
                or re.match(r"^\s*\d+\.\s+", nxt)
                or re.match(r"^\s*-{3,}\s*$", nxt)
            ):
                break
            chunk.append(clean_md(nxt))
            j += 1
        add_para(doc, " ".join(chunk), size=10.5, color=SLATE, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        i = j

    add_para(doc, "", space_after=2)
    add_para(
        doc,
        "© حقوق توسعه نرم‌افزار: شرکت وفاق شریف — طراح و توسعه‌دهنده اصلی: امید رامشینی — پشتیبانی فنی: 09398637969",
        bold=True,
        size=9,
        color=MUTED,
        space_after=2,
    )
    add_para(
        doc,
        "مخاطب: مدیرعامل / مدیریت ارشد پتروشیمی اروند — تاریخ تنظیم گزارش: ۱۴۰۵/۰۷/۱۱",
        size=9,
        color=MUTED,
        space_after=2,
    )

    try:
        doc.save(str(OUT))
        target = OUT
    except PermissionError:
        # File open in Word — write a sibling so the user can close & replace
        alt = OUT.with_name(OUT.stem + "-rtl" + OUT.suffix)
        doc.save(str(alt))
        target = alt
        print("WARN locked:", OUT.name, "-> wrote", alt.name)

    try:
        OUT_FA.write_bytes(target.read_bytes())
    except PermissionError:
        alt_fa = OUT_FA.with_name(OUT_FA.stem + "-rtl" + OUT_FA.suffix)
        alt_fa.write_bytes(target.read_bytes())
        print("WARN locked:", OUT_FA.name, "-> wrote", alt_fa.name)

    print("OK", target, "bytes=", target.stat().st_size)


if __name__ == "__main__":
    build()
