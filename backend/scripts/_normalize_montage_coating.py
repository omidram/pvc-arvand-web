"""Normalize «شرکت پوشش دهنده» / anode-cathode remarks for montage Excel import.

Reads:  d:\\مونتاژ - نصب و دی مونتاژ.xlsx
Writes: d:\\مونتاژ - نصب و دی مونتاژ - import-ready.xlsx
        backend/instance/montage_import_ready.xlsx
"""
from __future__ import annotations

import re
import shutil
from collections import Counter, defaultdict
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"d:\مونتاژ - نصب و دی مونتاژ.xlsx")
OUT_D = Path(r"d:\مونتاژ - نصب و دی مونتاژ - import-ready.xlsx")
OUT_INST = Path(__file__).resolve().parents[1] / "instance" / "montage_import_ready.xlsx"

COL_ANODE = 3
COL_CATHODE = 4
COL_COAT = 8
COL_ANODE_NOTE = 9
COL_CATHODE_NOTE = 10
COL_DAY = 12
COL_MONTH = 13
COL_YEAR = 14

DENORA = "دنورا"
DENORA_ALT = "دنوار"  # spelling used only for no-history dismantle fallback
PARTO = "پرتو صنعت پاژ"
ORIGINAL = "original(دنورا)"

_SPLIT_RE = re.compile(
    r"^(آند|کاتد)\s*(.+?)\s*(?:-|–|ـ|\s+و\s+)\s*(آند|کاتد)\s*(.+)$",
    re.I,
)
_SHARED_RE = re.compile(r"^آند\s*و\s*کاتد\s*(.+)$", re.I)
_SPACE_RE = re.compile(r"\s+")


def fold(text: str) -> str:
    text = text.replace("ي", "ی").replace("ك", "ک").replace("کاند", "کاتد")
    text = text.replace("وکاتد", "و کاتد").replace("وآند", "و آند")
    text = text.replace("دی\u200cمونتاژ", "دی مونتاژ").replace("دیمونتاژ", "دی مونتاژ")
    return _SPACE_RE.sub(" ", text).strip()


def clean_serial(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip().upper())


def jalali_tuple(day, month, year) -> tuple[int, int, int] | None:
    try:
        d, m, y = int(day), int(month), int(year)
    except (TypeError, ValueError):
        return None
    if y < 100:
        y = 1400 + y if y <= 80 else 1300 + y
    if not (1 <= m <= 12 and 1 <= d <= 31 and 1300 <= y <= 1500):
        return None
    return y, m, d


def append_note(existing, addition: str | None) -> str | None:
    addition = (addition or "").strip()
    if not addition:
        return existing if existing not in (None, "") else None
    cur = str(existing).strip() if existing not in (None, "") else ""
    if not cur:
        return addition
    if addition in cur:
        return cur
    return f"{cur} | {addition}"


def finalize_original(value: str | None) -> str | None:
    if value is None:
        return None
    text = fold(str(value))
    if not text:
        return value
    low = text.casefold()
    if low in {"original", "اورجینال", "اوریجینال"} or text == "original":
        return ORIGINAL
    # whole-field equals original variants already handled; also bare ORIGINAL(...)
    if re.fullmatch(r"original\s*(\(.*\))?", text, flags=re.I) and "دنورا" not in text and "دنوار" not in text:
        return ORIGINAL
    return text


def is_blank(text: str) -> bool:
    return not text.strip()


def is_xrf(text: str) -> bool:
    return bool(re.search(r"\bxrf\b", text, flags=re.I)) or "ایکس آر اف" in text


def is_purchased(text: str) -> bool:
    return "خریداری" in text


def is_repair(text: str) -> bool:
    return "تعمیر" in text


def is_dc_flange(text: str) -> bool:
    return ("لوز" in text or "فلنج" in text) and re.search(r"\bdc\b", text, flags=re.I)


def is_china(text: str) -> bool:
    return "چینی" in text or re.search(r"(?<![ا-ی])چین(?![ا-ی])", text) or text.strip() in {"چین", "چینی"}


def is_new(text: str) -> bool:
    t = text.strip()
    if t.upper() in {"NEW", "جدید"}:
        return True
    if "جدید" in t and "قابل" not in t:
        return True
    return False


def is_gen4(text: str) -> bool:
    return "نسل چهار" in text or "نسل چهارم" in text


def is_history_trigger(text: str) -> bool:
    if "تست بدنه" in text:
        return True
    if "قابل استفاده نو" in text:
        return True
    if "دی مونتاژ" in text or "دی‌مونتاژ" in text:
        return True
    return False


def is_already_original(text: str) -> bool:
    t = text.strip()
    low = t.casefold()
    return low in {"original", "اورجینال", "اوریجینال"} or low.startswith("original(")


def looks_like_denora(text: str) -> bool:
    return "دنورا" in text or "دنوار" in text


def split_sides(raw: str) -> tuple[str | None, str | None, str]:
    """Return (anode_part, cathode_part, mode) where mode is shared|split|single."""
    text = fold(raw)
    if not text:
        return "", "", "shared"
    m = _SPLIT_RE.match(text)
    if m and m.group(1) != m.group(3):
        parts = {m.group(1): m.group(2).strip(), m.group(3): m.group(4).strip()}
        return parts.get("آند", ""), parts.get("کاتد", ""), "split"
    m = _SHARED_RE.match(text)
    if m:
        rest = m.group(1).strip()
        return rest, rest, "shared"
    # "آند وکاتد ..." without space variants already folded
    if text.startswith("آند و کاتد") or text.startswith("آند وکاتد"):
        rest = re.sub(r"^آند\s*و\s*کاتد\s*", "", text).strip()
        return rest, rest, "shared"
    return text, text, "single"


def transform_side(
    phrase: str,
    *,
    serial: str,
    when: tuple[int, int, int] | None,
    history: dict[str, list[tuple[tuple[int, int, int], str, str | None]]],
    stats: Counter,
) -> tuple[str, str | None]:
    """Return (company, note_to_append)."""
    text = fold(phrase or "")
    if is_blank(text):
        stats["blank->original"] += 1
        return ORIGINAL, None

    if is_gen4(text):
        stats["gen4"] += 1
        return PARTO, text

    if is_xrf(text):
        stats["xrf"] += 1
        return DENORA, text

    if is_purchased(text):
        stats["purchased"] += 1
        return DENORA, text

    if is_repair(text):
        stats["repair"] += 1
        return DENORA, text

    if is_dc_flange(text):
        stats["dc_flange"] += 1
        return DENORA, text

    if is_china(text):
        stats["china"] += 1
        return DENORA, text

    if is_new(text):
        stats["new"] += 1
        return ORIGINAL, text

    if is_history_trigger(text):
        stats["history_trigger"] += 1
        hist = None
        if serial and when and serial in history:
            earlier = [item for item in history[serial] if item[0] < when]
            if earlier:
                hist = max(earlier, key=lambda item: item[0])
                hist = (hist[1], hist[2])
        if hist:
            stats["history_hit"] += 1
            company, note = hist
            note = append_note(note, text)
            return (finalize_original(company) or DENORA), note
        stats["history_miss"] += 1
        return DENORA_ALT, ORIGINAL

    if is_already_original(text):
        stats["original"] += 1
        return ORIGINAL, None

    if looks_like_denora(text):
        # e.g. «آند و کاتد دنورا» / «دنورا (قابل استفاده)»
        core = fold(re.sub(r"^(آند|کاتد)\s*", "", text))
        if core in {DENORA, DENORA_ALT} or re.fullmatch(rf"{DENORA}\s*(\(.*\))?", core):
            stats["denora_keep"] += 1
            return DENORA, text if core != DENORA else None
        stats["denora_wrap"] += 1
        return DENORA, text

    # Keep known plant vendors / other free text as coating company
    stats["keep"] += 1
    return text, None


def compose_coating(anode_co: str, cathode_co: str, mode: str) -> str:
    a = finalize_original(anode_co) or ""
    c = finalize_original(cathode_co) or ""
    if not a and not c:
        return ORIGINAL
    if a == c or mode == "shared" and a == c:
        return a or c
    if not a:
        return c
    if not c:
        return a
    return f"آند {a} - کاتد {c}"


def process_sheet(ws, stats: Counter) -> int:
    # Collect data rows then process in C90-date order so history lookups work.
    records: list[dict] = []
    for row_idx in range(3, ws.max_row + 1):
        anode = clean_serial(ws.cell(row_idx, COL_ANODE).value)
        cathode = clean_serial(ws.cell(row_idx, COL_CATHODE).value)
        raw_coat = ws.cell(row_idx, COL_COAT).value
        raw_text = "" if raw_coat is None else str(raw_coat)
        when = jalali_tuple(
            ws.cell(row_idx, COL_DAY).value,
            ws.cell(row_idx, COL_MONTH).value,
            ws.cell(row_idx, COL_YEAR).value,
        )
        # Skip empty trailing rows
        if not anode and not cathode and not raw_text.strip() and when is None:
            continue
        records.append(
            {
                "row": row_idx,
                "anode": anode,
                "cathode": cathode,
                "raw": raw_text,
                "when": when or (0, 0, 0),
                "anode_note": ws.cell(row_idx, COL_ANODE_NOTE).value,
                "cathode_note": ws.cell(row_idx, COL_CATHODE_NOTE).value,
            }
        )

    records.sort(key=lambda r: (r["when"], r["row"]))

    anode_hist: dict[str, list] = defaultdict(list)
    cathode_hist: dict[str, list] = defaultdict(list)
    changed = 0

    for rec in records:
        when = rec["when"] if rec["when"] != (0, 0, 0) else None
        a_part, c_part, mode = split_sides(rec["raw"])
        a_co, a_move = transform_side(
            a_part, serial=rec["anode"], when=when, history=anode_hist, stats=stats
        )
        c_co, c_move = transform_side(
            c_part, serial=rec["cathode"], when=when, history=cathode_hist, stats=stats
        )
        new_coat = compose_coating(a_co, c_co, mode)
        new_a_note = append_note(rec["anode_note"], a_move)
        new_c_note = append_note(rec["cathode_note"], c_move)
        if new_a_note and fold(str(new_a_note)) in {"original", "اورجینال", "اوریجینال"}:
            new_a_note = ORIGINAL
        if new_c_note and fold(str(new_c_note)) in {"original", "اورجینال", "اوریجینال"}:
            new_c_note = ORIGINAL

        row_idx = rec["row"]
        coat_cell = ws.cell(row_idx, COL_COAT)
        anode_note_cell = ws.cell(row_idx, COL_ANODE_NOTE)
        cathode_note_cell = ws.cell(row_idx, COL_CATHODE_NOTE)
        if (
            str(rec["raw"] or "") != str(new_coat or "")
            or str(rec["anode_note"] or "") != str(new_a_note or "")
            or str(rec["cathode_note"] or "") != str(new_c_note or "")
        ):
            changed += 1
        coat_cell.value = new_coat
        anode_note_cell.value = new_a_note
        cathode_note_cell.value = new_c_note

        if when:
            if rec["anode"]:
                anode_hist[rec["anode"]].append((when, a_co, new_a_note))
            if rec["cathode"]:
                cathode_hist[rec["cathode"]].append((when, c_co, new_c_note))

    return changed


def main() -> None:
    if not SRC.is_file():
        raise SystemExit(f"missing source: {SRC}")
    print(f"loading {SRC}", flush=True)
    wb = load_workbook(SRC)
    stats: Counter = Counter()
    sheet = wb.worksheets[0]
    print(f"sheet={sheet.title!r} rows={sheet.max_row}", flush=True)
    changed = process_sheet(sheet, stats)
    OUT_D.parent.mkdir(parents=True, exist_ok=True)
    OUT_INST.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT_D)
    shutil.copy2(OUT_D, OUT_INST)
    print(f"changed_rows={changed}", flush=True)
    for key, count in stats.most_common():
        print(f"  {key}={count}", flush=True)
    print(f"wrote {OUT_D}", flush=True)
    print(f"wrote {OUT_INST}", flush=True)


if __name__ == "__main__":
    main()
