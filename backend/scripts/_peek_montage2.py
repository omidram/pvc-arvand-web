"""Peek the latest montage workbook headers and status values."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

downloads = Path(r"C:\Users\Omid\Downloads")
matches = sorted(downloads.glob("*(2).xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
print("matches", len(matches))
path = matches[0]
print("file", path.name, path.stat().st_size)
wb = load_workbook(path, read_only=True, data_only=True)
print("sheets", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("\n=== SHEET", name)
    rows = []
    for i, row in enumerate(ws.iter_rows(max_row=8, values_only=True)):
        cells = []
        for c in row:
            if c is None:
                cells.append("")
            else:
                cells.append(str(c).replace("\n", " ")[:50])
        rows.append(cells)
        print(i, len(cells), " | ".join(cells[:45]))
wb.close()
