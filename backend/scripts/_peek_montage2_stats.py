"""Summarize status values and the repair-company sheet."""
from __future__ import annotations

from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

path = sorted(Path(r"C:\Users\Omid\Downloads").glob("*(2).xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)[0]
wb = load_workbook(path, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
status1 = Counter()
status2 = Counter()
use = Counter()
n = 0
out_with_date = 0
out_no_date = 0
in_service = 0
samples_other = []
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 2:
        continue
    n += 1
    s1 = "" if row[6] is None else str(row[6]).strip()
    s2 = "" if len(row) < 18 or row[17] is None else str(row[17]).strip()
    u = "" if row[4] is None else str(row[4]).strip()
    status1[s1] += 1
    status2[s2] += 1
    use[u] += 1
    dismantle = row[18] if len(row) > 18 else None
    if s2.upper() == "OUT OF SERVICE":
        if dismantle:
            out_with_date += 1
        else:
            out_no_date += 1
    elif s2.upper() == "IN SERVICE":
        in_service += 1
    elif s2 and len(samples_other) < 15:
        samples_other.append((s2, row[1], row[2], row[7], dismantle, row[19] if len(row) > 19 else None))

print("rows", n)
print("STATUS1")
for k, v in status1.most_common(30):
    print(f"  {v:6} {k!r}")
print("STATUS2")
for k, v in status2.most_common(30):
    print(f"  {v:6} {k!r}")
print("MEMBRANE USE")
for k, v in use.most_common(20):
    print(f"  {v:6} {k!r}")
print("out_with_date", out_with_date, "out_no_date", out_no_date, "in_service", in_service)
print("other status samples", samples_other)

ws3 = wb[wb.sheetnames[2]]
print("\n=== sheet3", wb.sheetnames[2])
c = 0
for i, row in enumerate(ws3.iter_rows(values_only=True)):
    c += 1
    if i < 12:
        print(i, row)
print("sheet3 rows", c)
wb.close()
