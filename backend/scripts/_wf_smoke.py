"""End-to-end run of both warehouse workflows on a temp SQLite DB (C-90 / procurement / warehouse users)."""
import io
import os
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from sqlalchemy import create_engine  # noqa: E402
from sqlalchemy.orm import sessionmaker  # noqa: E402

from app import models, workflow_models  # noqa: E402,F401
from app.auth import get_current_user, seed_default_roles  # noqa: E402
from app.database import get_db  # noqa: E402
from app.routers import workflow  # noqa: E402

tmp = Path(tempfile.mkdtemp(prefix="wf_smoke_"))
workflow.data_root = lambda: tmp
engine = create_engine(f"sqlite:///{tmp / 'wf.db'}", connect_args={"check_same_thread": False})
models.Base.metadata.create_all(engine)
Session = sessionmaker(bind=engine)

with Session() as db:
    seed_default_roles(db)
    roles = {r.name: r.id for r in db.query(models.AppRole).all()}
    for name, role in (("c90", "C-90"), ("proc", "Procurement"), ("wh", "Warehouse"), ("viewer", "Viewer")):
        db.add(models.User(username=name, full_name=name.upper(), password_hash="x", role="user", role_id=roles[role]))
    db.add(models.WhCompany(name="CoatCo", is_coating=True, is_supplier=False))
    db.add(models.WhCompany(name="SupCo", is_coating=False, is_supplier=True))
    for i in range(1, 46):
        db.add(models.Anode(anode_nr=f"A-{i:03d}"))
    for i in range(1, 6):
        db.add(models.Cathode(cathode_nr=f"K-{i:03d}"))
    db.commit()
    users = {u.username: u.id for u in db.query(models.User).all()}

app = FastAPI()
app.include_router(workflow.router)
current = {"name": "c90"}


def _db():
    db = Session()
    try:
        yield db
    finally:
        db.close()


def _user():
    with Session() as db:
        u = db.get(models.User, users[current["name"]])
        db.expunge(u)
        return u


app.dependency_overrides[get_db] = _db
app.dependency_overrides[get_current_user] = _user
client = TestClient(app)


def as_(name):
    current["name"] = name
    return client


def ok(resp, code=200):
    assert resp.status_code == code, (resp.status_code, resp.text)
    return resp.json() if resp.headers.get("content-type", "").startswith("application/json") else resp


# ---------------- recoating
r = as_("c90").post("/workflow/coating", json={"anodes": [{"serial": "A-999"}], "cathodes": []})
assert r.status_code == 422 and r.json()["detail"]["code"] == "unknown_serials", r.text
assert as_("proc").post("/workflow/coating", json={"anodes": [{"serial": "A-001"}]}).status_code == 403

anodes = [{"serial": f"a-{i:03d}", "problem": "dent" if i == 2 else None} for i in range(1, 46)]
d = ok(as_("c90").post("/workflow/coating", json={"anodes": anodes, "cathodes": [{"serial": "K-001"}, {"serial": "K-002"}]}))
codes = d["codes"]
print("codes:", codes)
assert len(codes) == 4 and codes[0].startswith("CO-") and codes[0].endswith("-0001") and codes[-1].endswith("-0004")
assert [len(p["items"]) for p in d["packages"]] == [20, 20, 5, 2]
assert d["packages"][0]["items"][0]["serial"] == "A-001"  # mapped to the catalog spelling
bid = d["id"]
assert as_("c90").post("/workflow/coating", json={"anodes": [{"serial": "A-001"}]}).status_code == 409

s = ok(as_("proc").get("/workflow/summary"))
assert s["cartable"] == 1 and s["unread"] == 1, s
assert ok(as_("wh").get("/workflow/summary"))["cartable"] == 0

as_("wh")
assert client.post(f"/workflow/coating/{bid}/company", json={"company_id": 1}).status_code == 403
with Session() as db:
    coat_id = db.query(models.WhCompany).filter_by(name="CoatCo").one().id
    sup_id = db.query(models.WhCompany).filter_by(name="SupCo").one().id
assert as_("proc").post(f"/workflow/coating/{bid}/company", json={"company_id": sup_id}).status_code == 400
d = ok(as_("proc").post(f"/workflow/coating/{bid}/company", json={"company_id": coat_id}))
assert d["status"] == "await_indent"
assert ok(as_("wh").get("/workflow/summary"))["unread"] == 1
d = ok(as_("wh").post(f"/workflow/coating/{bid}/indent", json={"indent_no": "IND-77"}))
assert d["status"] == "await_dispatch"
d = ok(as_("proc").post(f"/workflow/coating/{bid}/dispatch", json={"dispatch_date": "2026-01-01", "expected_return_date": "2026-01-21"}))
assert d["status"] == "at_coater"
r = as_("wh").post(f"/workflow/coating/{bid}/receive", json={"actual_date": "2026-02-10"})
assert r.status_code == 400, r.text  # late without a reason
d = ok(as_("wh").post(f"/workflow/coating/{bid}/receive", json={"actual_date": "2026-02-10", "delay_party": "company", "delay_reason": "Late furnace"}))
assert d["status"] == "await_report" and d["delay_days"] == 20

# Excel template -> fill -> upload
x = ok(as_("proc").get(f"/workflow/coating/{bid}/excel"))
wb = load_workbook(io.BytesIO(x.content))
assert wb.sheetnames == codes, wb.sheetnames
ws = wb[codes[0]]
hdr = next(i for i, row in enumerate(ws.iter_rows(values_only=True), start=1) if row and row[2] == "Element No.")
ws.cell(row=hdr + 1, column=5, value="pitting")
ws.cell(row=hdr + 1, column=6, value="recoated OK")
ws.cell(row=hdr + 2, column=7, value="edge repaired")
buf = io.BytesIO()
wb.save(buf)
assert as_("proc").post(f"/workflow/coating/{bid}/submit", json={}).status_code == 400  # no report yet
res = ok(as_("proc").post(f"/workflow/coating/{bid}/report", files={"files": ("filled.xlsx", buf.getvalue())}))
assert res["rows_filled"] == 47, res["rows_filled"]
item = res["detail"]["packages"][0]["items"][0]
assert item["final_report"] == "recoated OK" and item["pre_inspection"] == "pitting"
d = ok(as_("proc").post(f"/workflow/coating/{bid}/submit", json={"note": "report attached"}))
assert d["status"] == "await_approval"
assert as_("c90").post(f"/workflow/coating/{bid}/reject", json={}).status_code == 400
d = ok(as_("c90").post(f"/workflow/coating/{bid}/reject", json={"note": "K-002 not coated"}))
assert d["status"] == "await_report" and d["rejected"]
d = ok(as_("proc").post(f"/workflow/coating/{bid}/rework", json={"dispatch_date": "2026-02-12", "expected_return_date": "2026-02-20"}))
assert d["status"] == "at_coater" and d["rework_count"] == 1
d = ok(as_("wh").post(f"/workflow/coating/{bid}/receive", json={"actual_date": "2026-02-18"}))
assert d["delay_days"] == -2
d = ok(as_("proc").post(f"/workflow/coating/{bid}/submit", json={}))
d = ok(as_("c90").post(f"/workflow/coating/{bid}/approve", json={"note": "ok"}))
assert d["status"] == "completed" and "approve" not in d["actions"]
with Session() as db:
    rec = db.query(models.AnodeRecoating).filter_by(anode_nr="A-001").one()
    assert rec.coating_nr == codes[0] and rec.manufacturer == "CoatCo" and rec.recoating_number == "1"
    st = db.query(models.WhElementState).filter_by(element_kind="anode", serial="A-001").one()
    assert st.lifecycle_status == "in_plant"

# ---------------- purchase
assert as_("c90").post("/workflow/purchase", json={"supplier_id": sup_id, "anode_count": 2}).status_code == 403
p = ok(as_("proc").post("/workflow/purchase", json={"supplier_id": sup_id, "anode_count": 2, "cathode_count": 1, "approval_ref": "MGT-5"}))
pid = p["id"]
assert p["code"].startswith("PO-") and p["code"].endswith("-0001") and p["status"] == "await_indent"
ok(as_("proc").post(f"/workflow/{'purchase'}/{pid}/files", files={"files": ("approval.pdf", b"%PDF-1.4")}))
p = ok(as_("wh").post(f"/workflow/purchase/{pid}/indent", json={"indent_no": "IND-88"}))
p = ok(as_("proc").post(f"/workflow/purchase/{pid}/order", json={"request_date": "2026-03-01", "expected_delivery_date": "2026-03-10"}))
p = ok(as_("wh").post(f"/workflow/purchase/{pid}/deliver", json={"actual_date": "2026-03-09"}))
assert p["status"] == "await_registration" and p["delay_days"] == -1
r = as_("c90").post(f"/workflow/purchase/{pid}/register", json={"items": [{"element_kind": "anode", "serial": "A-001"}]})
assert r.status_code == 409, r.text
p = ok(as_("c90").post(f"/workflow/purchase/{pid}/register", json={"items": [{"element_kind": "anode", "serial": "NA-1", "coating": "RuO2"}]}))
assert p["status"] == "await_registration" and p["registered"] == 1
r = as_("c90").post(f"/workflow/purchase/{pid}/register", json={"items": [{"element_kind": "anode", "serial": "NA-2"}, {"element_kind": "anode", "serial": "NA-3"}]})
assert r.status_code == 400, r.text
p = ok(as_("c90").post(f"/workflow/purchase/{pid}/register", json={"items": [{"element_kind": "anode", "serial": "NA-2"}, {"element_kind": "cathode", "serial": "NK-1"}]}))
assert p["status"] == "completed"
with Session() as db:
    a = db.get(models.Anode, "NA-1")
    assert a.manufacturer == "SupCo" and a.coating == "RuO2" and a.batch == p["code"]

# ---------------- history / reports / notifications
h = ok(as_("viewer").get("/workflow/history", params={"serial": "A-001"}))
assert h["rows"] and h["rows"][0]["code"] == codes[0] and h["rows"][0]["final_report"] == "recoated OK", h["rows"][:1]
h = ok(as_("viewer").get("/workflow/history", params={"code": codes[3]}))
assert {r["serial"] for r in h["rows"]} == {"K-001", "K-002"}
h = ok(as_("viewer").get("/workflow/history", params={"date_from": "2026-03-01"}))
assert h["totals"]["purchases"] == 1 and h["totals"]["coating_pieces"] == 0, h["totals"]
rep = {r["company_name"]: r for r in ok(as_("viewer").get("/workflow/reports/companies"))}
assert rep["CoatCo"]["coating_pieces"] == 47 and rep["CoatCo"]["reworks"] == 1
assert rep["SupCo"]["purchases"] == 1 and rep["SupCo"]["on_time"] == 1
assert as_("viewer").get("/workflow/cartable").json() == []
n = ok(as_("proc").get("/workflow/notifications"))
assert n, "procurement got no notifications"
ok(as_("proc").post("/workflow/notifications/read", json={"all": True}))
assert ok(as_("proc").get("/workflow/summary"))["unread"] == 0
fid = ok(as_("viewer").get(f"/workflow/purchase/{pid}"))["files"][0]["id"]
assert as_("viewer").get(f"/workflow/files/{fid}").content == b"%PDF-1.4"
print("ALL WORKFLOW CHECKS PASSED")
