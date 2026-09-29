"""Load anode and cathode serial lists into the component catalogs.

Existing detail rows are kept. A missing serial is added. An empty remarks
field is filled from the workbook note. Nothing else is overwritten.
"""
from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

BACKEND = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND))

from app.database import SessionLocal  # noqa: E402
from app import models  # noqa: E402
from app.warehouse import compact_nr  # noqa: E402

FILES = [
    (
        Path(r"C:\Users\Omid\Downloads\ANODE  SERIAL NUMBERS.xlsx"),
        "ANODE",
        models.Anode,
        "anode_nr",
    ),
    (
        Path(r"C:\Users\Omid\Downloads\CATHODE SERIAL NUMBERS.xlsx"),
        "CATHODE",
        models.Cathode,
        "cathode_nr",
    ),
]


def rows(path: Path, sheet: str):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet]
    out = []
    for index, row in enumerate(ws.iter_rows(values_only=True)):
        if index == 0 or not row:
            continue
        serial = row[1] if len(row) > 1 else None
        note = row[2] if len(row) > 2 else None
        text = "" if serial is None else str(serial).strip()
        if text.endswith(".0") and text[:-2].isdigit():
            text = text[:-2]
        remark = "" if note is None else str(note).strip()
        if text:
            out.append((text[:50], remark[:500] if remark else None))
    wb.close()
    return out


def main() -> None:
    with SessionLocal() as db:
        for path, sheet, model, field in FILES:
            loaded = rows(path, sheet)
            existing = {
                compact_nr(getattr(row, field)): row
                for row in db.query(model).all()
                if compact_nr(getattr(row, field))
            }
            created = 0
            noted = 0
            present = 0
            seen: set[str] = set()
            for serial, remark in loaded:
                key = compact_nr(serial)
                if not key or key in seen:
                    continue
                seen.add(key)
                row = existing.get(key)
                if row is None:
                    db.add(model(**{field: serial, "remarks": remark}))
                    existing[key] = True
                    created += 1
                    continue
                present += 1
                if remark and not (getattr(row, "remarks", None) or "").strip():
                    row.remarks = remark
                    noted += 1
            db.commit()
            total = db.query(model).count()
            print(f"{sheet}: file {len(seen)} created {created} already {present} remarks filled {noted} catalog {total}")


if __name__ == "__main__":
    main()
