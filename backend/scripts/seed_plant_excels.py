"""Load plant Excel catalogues into the live SQLite DB.

Sources (repo copies under docs/plant-data/):
  - ANODE-SERIAL-NUMBERS.xlsx  -> anodes
  - CATHODE-SERIAL-NUMBERS.xlsx -> cathodes
  - montage-install-dismantle-import-ready.xlsx -> elements (assembly)
  - TAFKIK.xlsx -> electrode_segregations
"""
from __future__ import annotations

import re
import sys
from pathlib import Path

import openpyxl
from sqlalchemy import text

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "backend"))

from app import models  # noqa: E402
from app.database import SessionLocal  # noqa: E402
from app.plant_import import ensure_electrode_serial_prefix, parse_assembly_excel, parse_tafkik_excel  # noqa: E402
from app.routers.elements import _apply_assembly_import  # noqa: E402
from app.segregation_enrich import build_latest_element_index, enrich_segregation_fields  # noqa: E402
from app.warehouse import compact_nr  # noqa: E402

DATA = ROOT / "docs" / "plant-data"


def _is_digits_only(value: str | None) -> bool:
    if not value:
        return False
    compact = re.sub(r"\s+", "", str(value))
    if any(ch.isalpha() for ch in compact):
        return False
    digits = compact[:-2] if compact.endswith(".0") and compact[:-2].isdigit() else compact.replace(".", "", 1)
    return digits.isdigit()


def migrate_digit_only_serials(db) -> dict:
    """Rename bare numeric anode/cathode PKs to A… / C… and rewrite FK columns."""
    anode_map: dict[str, str] = {}
    for (nr,) in db.query(models.Anode.anode_nr).all():
        if not _is_digits_only(nr):
            continue
        target = ensure_electrode_serial_prefix(nr, "anode")
        if not target or target == nr:
            continue
        anode_map[nr] = target
        if not db.get(models.Anode, target):
            db.execute(
                text("UPDATE anodes SET anode_nr = :new WHERE anode_nr = :old"),
                {"old": nr, "new": target},
            )

    cathode_map: dict[str, str] = {}
    for (nr,) in db.query(models.Cathode.cathode_nr).all():
        if not _is_digits_only(nr):
            continue
        target = ensure_electrode_serial_prefix(nr, "cathode")
        if not target or target == nr:
            continue
        cathode_map[nr] = target
        if not db.get(models.Cathode, target):
            db.execute(
                text("UPDATE cathodes SET cathode_nr = :new WHERE cathode_nr = :old"),
                {"old": nr, "new": target},
            )

    db.expire_all()

    def rewrite(table: str, column: str, mapping: dict[str, str], kind: str) -> int:
        changed = 0
        for old, new in mapping.items():
            res = db.execute(
                text(f"UPDATE {table} SET {column} = :new WHERE {column} = :old"),
                {"old": old, "new": new},
            )
            changed += res.rowcount or 0
        rows = db.execute(text(f"SELECT DISTINCT {column} FROM {table} WHERE {column} IS NOT NULL")).fetchall()
        for (value,) in rows:
            if not _is_digits_only(value):
                continue
            new = ensure_electrode_serial_prefix(value, kind)
            if not new or new == value:
                continue
            res = db.execute(
                text(f"UPDATE {table} SET {column} = :new WHERE {column} = :old"),
                {"old": value, "new": new},
            )
            changed += res.rowcount or 0
        return changed

    stats = {
        "anodes_renamed": len(anode_map),
        "cathodes_renamed": len(cathode_map),
        "elements_anode": rewrite("elements", "anode_nr", anode_map, "anode"),
        "elements_cathode": rewrite("elements", "cathode_nr", cathode_map, "cathode"),
        "anode_maintenance": rewrite("anode_maintenance", "anode_nr", anode_map, "anode"),
        "cathode_maintenance": rewrite("cathode_maintenance", "cathode_nr", cathode_map, "cathode"),
        "anode_recoating": rewrite("anode_recoating", "anode_nr", anode_map, "anode"),
        "cathode_recoating": rewrite("cathode_recoating", "cathode_nr", cathode_map, "cathode"),
        "anode_coating_checks": rewrite("anode_coating_checks", "anode_nr", anode_map, "anode"),
        "cathode_coating_checks": rewrite("cathode_coating_checks", "cathode_nr", cathode_map, "cathode"),
    }

    for table in ("inspection_reports", "assembly_inspection_reports"):
        exists = db.execute(
            text("SELECT 1 FROM sqlite_master WHERE type='table' AND name=:n"),
            {"n": table},
        ).fetchone()
        if not exists:
            continue
        stats[f"{table}_anode"] = rewrite(table, "anode_nr", anode_map, "anode")
        stats[f"{table}_cathode"] = rewrite(table, "cathode_nr", cathode_map, "cathode")
    anode_targets = {compact_nr(v) for v in anode_map.values()} | {compact_nr(k) for k in anode_map}
    cathode_targets = {compact_nr(v) for v in cathode_map.values()} | {compact_nr(k) for k in cathode_map}
    seg_changed = 0
    for row in db.query(models.ElectrodeSegregation).all():
        kind = row.electrode_kind
        if kind not in ("anode", "cathode"):
            key = compact_nr(row.serial_nr)
            if key in anode_targets or f"A{key}" in anode_targets:
                kind = "anode"
            elif key in cathode_targets or f"C{key}" in cathode_targets:
                kind = "cathode"
        if kind in ("anode", "cathode") and _is_digits_only(row.serial_nr):
            new = ensure_electrode_serial_prefix(row.serial_nr, kind)
            if new and new != row.serial_nr:
                row.serial_nr = new
                row.electrode_kind = kind
                seg_changed += 1
        if row.pair_serial_nr and _is_digits_only(row.pair_serial_nr):
            pair_kind = "cathode" if (kind or row.electrode_kind) == "anode" else "anode"
            new_pair = ensure_electrode_serial_prefix(row.pair_serial_nr, pair_kind)
            if new_pair and new_pair != row.pair_serial_nr:
                row.pair_serial_nr = new_pair
                seg_changed += 1

    for old, new in anode_map.items():
        if old != new:
            db.execute(text("DELETE FROM anodes WHERE anode_nr = :old"), {"old": old})
    for old, new in cathode_map.items():
        if old != new:
            db.execute(text("DELETE FROM cathodes WHERE cathode_nr = :old"), {"old": old})

    stats["segregation"] = seg_changed
    db.commit()
    return stats

def _norm_serial(value, *, kind: str | None = None) -> str:
    if value is None:
        return ""
    text = " ".join(str(value).split()).strip()
    if kind in ("anode", "cathode"):
        prefixed = ensure_electrode_serial_prefix(text, kind)
        return prefixed or ""
    return text


def seed_anode_serials(db) -> int:
    path = DATA / "ANODE-SERIAL-NUMBERS.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["ANODE"] if "ANODE" in wb.sheetnames else wb[wb.sheetnames[0]]
    created = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        serial = _norm_serial(row[1] if len(row) > 1 else None, kind="anode")
        if not serial:
            continue
        remark = _norm_serial(row[2] if len(row) > 2 else None) or None
        generation = None
        if remark and "نسل" in remark:
            generation = remark
        existing = db.get(models.Anode, serial)
        if existing:
            if remark and not existing.remarks:
                existing.remarks = remark
            if generation and not existing.generation:
                existing.generation = generation
            continue
        db.add(models.Anode(anode_nr=serial, remarks=remark, generation=generation))
        created += 1
    wb.close()
    db.commit()
    return created


def seed_cathode_serials(db) -> int:
    path = DATA / "CATHODE-SERIAL-NUMBERS.xlsx"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["CATHODE"] if "CATHODE" in wb.sheetnames else wb[wb.sheetnames[0]]
    created = 0
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        serial = _norm_serial(row[1] if len(row) > 1 else None, kind="cathode")
        if not serial:
            continue
        remark = _norm_serial(row[2] if len(row) > 2 else None) or None
        generation = None
        if remark and "نسل" in remark:
            generation = remark
        existing = db.get(models.Cathode, serial)
        if existing:
            if remark and not existing.remarks:
                existing.remarks = remark
            if generation and not existing.generation:
                existing.generation = generation
            continue
        db.add(models.Cathode(cathode_nr=serial, remarks=remark, generation=generation))
        created += 1
    wb.close()
    db.commit()
    return created


def seed_montage(db) -> dict:
    path = DATA / "montage-install-dismantle-import-ready.xlsx"
    content = path.read_bytes()
    parsed = parse_assembly_excel(content)
    if not parsed.get("recognized"):
        raise SystemExit(f"Montage workbook not recognized: {path}")
    return _apply_assembly_import(db, parsed, mode="upsert", progress=None)


def _seg_kwargs(fields: dict) -> dict:
    from datetime import date, datetime

    names = {c.name for c in models.ElectrodeSegregation.__table__.columns}
    out = {}
    for key, value in fields.items():
        if key not in names:
            continue
        if key.endswith("_date"):
            if value is None or value == "":
                out[key] = None
            elif isinstance(value, date) and not isinstance(value, datetime):
                out[key] = value
            else:
                try:
                    out[key] = date.fromisoformat(str(value)[:10])
                except ValueError:
                    out[key] = None
        else:
            out[key] = value
    return out


def seed_tafkik(db) -> int:
    path = DATA / "TAFKIK.xlsx"
    content = path.read_bytes()
    parsed = parse_tafkik_excel(content)
    records = parsed.get("records") or []
    anode_keys = {compact_nr(nr) for (nr,) in db.query(models.Anode.anode_nr).all() if nr}
    cathode_keys = {compact_nr(nr) for (nr,) in db.query(models.Cathode.cathode_nr).all() if nr}
    element_index = build_latest_element_index(db)
    existing = {
        (compact_nr(s), str(d) if d else "")
        for s, d in db.query(
            models.ElectrodeSegregation.serial_nr,
            models.ElectrodeSegregation.inspection_date,
        ).all()
        if s
    }
    created = 0
    for item in records:
        kind = item.get("electrode_kind")
        raw_serial = item["serial_nr"]
        if not kind or kind == "unknown":
            key = compact_nr(raw_serial)
            # Match digit-only Excel rows to prefixed master lists (A123 / C123).
            if key in anode_keys and key not in cathode_keys:
                kind = "anode"
            elif key in cathode_keys and key not in anode_keys:
                kind = "cathode"
            elif f"A{key}" in anode_keys and f"C{key}" not in cathode_keys:
                kind = "anode"
            elif f"C{key}" in cathode_keys and f"A{key}" not in anode_keys:
                kind = "cathode"
        if kind in ("anode", "cathode"):
            serial = ensure_electrode_serial_prefix(raw_serial, kind) or raw_serial
            pair = item.get("pair_serial_nr")
            pair_kind = "cathode" if kind == "anode" else "anode"
            if pair:
                item = {**item, "pair_serial_nr": ensure_electrode_serial_prefix(pair, pair_kind) or pair}
        else:
            serial = raw_serial
        identity = (compact_nr(serial), str(item.get("inspection_date") or ""))
        if identity in existing:
            continue
        fields = enrich_segregation_fields(
            db,
            serial_nr=serial,
            electrode_kind=kind,
            preserve_service_life=item.get("service_life") or None,
            excel_values={**item, "serial_nr": serial},
            element_index=element_index,
        )
        row = models.ElectrodeSegregation(**_seg_kwargs(fields))
        db.add(row)
        existing.add(identity)
        created += 1
        if created % 500 == 0:
            db.flush()
    db.commit()
    return created


def main() -> None:
    missing = [
        name
        for name in (
            "ANODE-SERIAL-NUMBERS.xlsx",
            "CATHODE-SERIAL-NUMBERS.xlsx",
            "montage-install-dismantle-import-ready.xlsx",
            "TAFKIK.xlsx",
        )
        if not (DATA / name).is_file()
    ]
    if missing:
        raise SystemExit(f"Missing plant-data files: {missing}")

    db = SessionLocal()
    try:
        a = seed_anode_serials(db)
        c = seed_cathode_serials(db)
        migrated = migrate_digit_only_serials(db)
        m = seed_montage(db)
        t = seed_tafkik(db)
        print(
            "OK anodes_new=",
            a,
            "cathodes_new=",
            c,
            "migrated=",
            migrated,
            "montage=",
            m,
            "tafkik_new=",
            t,
            "anode_total=",
            db.query(models.Anode).count(),
            "cathode_total=",
            db.query(models.Cathode).count(),
            "element_total=",
            db.query(models.Element).count(),
            "segregation_total=",
            db.query(models.ElectrodeSegregation).count(),
        )
    finally:
        db.close()


if __name__ == "__main__":
    main()
